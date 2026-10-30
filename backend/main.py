from fastapi import FastAPI, Depends, HTTPException, Query, BackgroundTasks, UploadFile, File, Request
from fastapi.middleware.cors import CORSMiddleware
from sqlalchemy.orm import Session
from sqlalchemy import func, or_, desc, case
from datetime import datetime, timedelta, timezone
from zoneinfo import ZoneInfo
from typing import Optional, List
from ai_service import get_ai_explanation, check_rate_limit
from apscheduler.schedulers.background import BackgroundScheduler
import random
import os
import secrets
import csv
import io
import openpyxl
import httpx
import hmac
import hashlib
from dotenv import load_dotenv

load_dotenv()

from database import get_db, Base, engine
import models
import email_service
from auth import hash_password, verify_password, create_access_token, generate_username
from schemas import *
from deps import get_current_user, require_admin, require_super_admin, is_school_scoped
from constants import (
    SUBJECTS_BY_LEVEL, EXAM_TARGETS_BY_LEVEL, EDUCATION_LEVELS,
    ALL_EXAM_TARGETS, ALL_SUBJECTS,
)

from models import (
    DailyChallengeLog, Question, UserInventory, Transaction, StudyMaterial,
    Squad, SquadMember, ReportedQuestion, StagingQuestion, Passage, StudySession
)

Base.metadata.create_all(bind=engine)

app = FastAPI(title="Makatron CBT API", version="0.1.0")

import os

# ... (other imports)

app = FastAPI(title="Makatron CBT API", version="0.1.0")

# Get the frontend URL from Render environment variables
frontend_url = os.getenv("FRONTEND_URL", "http://localhost:5173")

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "http://127.0.0.1:5173",
        "http://localhost:3000",
        frontend_url,
    ],
    allow_origin_regex=r"https://.*\.netlify\.app|https://.*\.vercel\.app", # Allows any Netlify/Vercel link
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

LAGOS = ZoneInfo("Africa/Lagos")


def week_start_lagos() -> datetime:
    now = datetime.now(LAGOS)
    monday = now - timedelta(days=now.weekday())
    monday = monday.replace(hour=0, minute=0, second=0, microsecond=0)
    return monday.astimezone(timezone.utc)


def log_activity(db: Session, action: str, description: str):
    db.add(models.ActivityLog(action_type=action, description=description))
    db.commit()


def _abandon_stale_tests(db: Session, user_id: int):
    now = datetime.now(timezone.utc)
    stale = db.query(models.Test).filter(
        models.Test.user_id == user_id,
        models.Test.status == "in_progress",
    ).all()
    for t in stale:
        if t.started_at and t.started_at + timedelta(seconds=t.duration_seconds) < now:
            t.status = "abandoned"
            t.submitted_at = t.started_at + timedelta(seconds=t.duration_seconds)
    db.commit()


def check_trial_access(user: models.User):
    if user.is_premium:
        if user.premium_expires_at and user.premium_expires_at < datetime.now(timezone.utc):
            raise HTTPException(403, "Your premium subscription has expired. Please renew.")
        return True
    if user.trial_ends_at and user.trial_ends_at < datetime.now(timezone.utc):
        raise HTTPException(403, "Your 7-day free trial has ended. Please upgrade to Premium.")
    return True


def require_premium(user: models.User):
    if not user.is_premium:
        raise HTTPException(403, "This feature is Premium-only. Please upgrade.")
    if user.premium_expires_at and user.premium_expires_at < datetime.now(timezone.utc):
        raise HTTPException(403, "Your premium subscription has expired. Please renew.")
    return True


def get_passage_payload(db: Session, passage_id: Optional[int]):
    if not passage_id:
        return None
    p = db.query(models.Passage).filter(models.Passage.id == passage_id).first()
    if not p:
        return None
    return {"title": p.title, "content": p.content}


def _touch_streak(user: models.User, db: Session):
    today = datetime.now(timezone.utc).date()
    if user.last_login_date == today:
        return
    if user.last_login_date is None:
        user.current_streak = 1
    else:
        gap_days = (today - user.last_login_date).days
        if gap_days == 1:
            user.current_streak = (user.current_streak or 0) + 1
        elif gap_days > 1:
            user.current_streak = 1
    user.last_login_date = today
    db.commit()


def _scoped_user_or_403(db: Session, admin: models.User, target_user_id: int) -> models.User:
    target = db.query(models.User).filter(models.User.id == target_user_id).first()
    if not target:
        raise HTTPException(404, "User not found")
    if is_school_scoped(admin) and target.school_id != admin.school_id:
        raise HTTPException(403, "You can only manage users within your school.")
    return target


# ---------------- SCHEDULED NOTIFICATIONS ----------------

def cron_send_notifications():
    from database import SessionLocal
    db = SessionLocal()
    try:
        now = datetime.now(timezone.utc)
        tomorrow = now + timedelta(days=1)

        upcoming = db.query(models.Exam).filter(
            models.Exam.exam_type == "championship",
            models.Exam.is_active == True,
            models.Exam.start_date >= now,
            models.Exam.start_date <= tomorrow,
        ).all()
        for exam in upcoming:
            students = db.query(models.User).filter(
                models.User.education_level == exam.education_level,
                models.User.is_admin == False,
                models.User.email != None,
            ).all()
            for s in students:
                try:
                    email_service.send_support_email(
                        s.display_name, s.email,
                        f"🏆 The {exam.title} starts soon on MAKAVERSE. Get ready to compete!"
                    )
                except Exception:
                    pass

        expiring = db.query(models.User).filter(
            models.User.trial_ends_at >= now,
            models.User.trial_ends_at <= tomorrow,
            models.User.is_premium == False,
            models.User.email != None,
        ).all()
        for u in expiring:
            try:
                email_service.send_support_email(
                    u.display_name, u.email,
                    "⏳ Your 7-day MAKAVERSE Premium trial ends tomorrow. Upgrade to keep full access."
                )
            except Exception:
                pass
    except Exception as e:
        print(f"Cron notification error: {e}")
    finally:
        db.close()


scheduler = BackgroundScheduler()
scheduler.add_job(cron_send_notifications, "cron", hour=8, minute=0)
scheduler.start()


@app.on_event("shutdown")
def shutdown_scheduler():
    scheduler.shutdown()


# ---------------- ROOT ----------------

@app.get("/")
def root(): return {"name": "Makatron CBT API", "status": "ok"}


@app.get("/health")
def health(): return {"status": "healthy"}


# ---------------- SCHOOLS ----------------

@app.get("/schools")
def get_public_schools(db: Session = Depends(get_db)):
    schools = db.query(models.School).order_by(models.School.name.asc()).all()
    return [{"id": s.id, "name": s.name, "state": s.state} for s in schools]


@app.post("/admin/schools")
def admin_create_school(payload: dict, db: Session = Depends(get_db), admin: models.User = Depends(require_super_admin)):
    name = payload.get("name", "").strip()
    state = payload.get("state", "").strip()
    if not name: raise HTTPException(400, "School name is required")
    existing = db.query(models.School).filter(func.lower(models.School.name) == name.lower()).first()
    if existing: raise HTTPException(400, "School already exists")
    school = models.School(name=name, state=state or None)
    db.add(school); db.commit(); db.refresh(school)
    log_activity(db, "SCHOOL_CREATED", f"Admin added school: {name}")
    return {"ok": True, "id": school.id}


@app.post("/admin/schools/upload")
async def admin_upload_schools(
    file: UploadFile = File(...),
    db: Session = Depends(get_db), admin: models.User = Depends(require_super_admin)
):
    filename = file.filename.lower()
    if not (filename.endswith(".csv") or filename.endswith(".xlsx")):
        raise HTTPException(400, "File must be .csv or .xlsx")
    contents = await file.read()
    rows_data = []
    if filename.endswith(".csv"):
        try: text = contents.decode("utf-8-sig")
        except: text = contents.decode("latin-1")
        reader = csv.DictReader(io.StringIO(text))
        for r in reader: rows_data.append(r)
    else:
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
        sheet = wb.active
        headers = [str(cell.value).strip().lower() if cell.value else "" for cell in sheet[1]]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not any(row): continue
            row_dict = {headers[i]: str(row[i]).strip() if row[i] is not None else "" for i in range(len(headers)) if i < len(row)}
            rows_data.append(row_dict)
    added = 0
    errors = []
    for row_num, row in enumerate(rows_data, start=2):
        try:
            name = row.get("name", "").strip()
            state = row.get("state", "").strip()
            if not name: raise ValueError("Missing school name")
            existing = db.query(models.School).filter(func.lower(models.School.name) == name.lower()).first()
            if existing: continue
            school = models.School(name=name, state=state or None)
            db.add(school); added += 1
        except Exception as e:
            errors.append(f"Row {row_num}: {str(e)}")
    db.commit()
    log_activity(db, "SCHOOLS_UPLOADED", f"Admin uploaded {added} schools")
    return {"added": added, "errors": errors}


@app.delete("/admin/schools/{school_id}")
def admin_delete_school(school_id: int, db: Session = Depends(get_db), admin: models.User = Depends(require_super_admin)):
    school = db.query(models.School).filter(models.School.id == school_id).first()
    if not school: raise HTTPException(404, "School not found")
    db.query(models.User).filter(models.User.school_id == school_id).update({"school_id": None})
    db.delete(school); db.commit()
    return {"ok": True}


# ---------------- AUTH ----------------

@app.post("/auth/register", response_model=RegisterResponse)
def register(payload: RegisterRequest, background_tasks: BackgroundTasks, db: Session = Depends(get_db)):
    settings = db.query(models.Settings).first()
    if settings and not settings.registration_open:
        raise HTTPException(400, "Registration is currently closed.")

    username = generate_username()
    while db.query(models.User).filter(models.User.username == username).first():
        username = generate_username()

    final_school_id = getattr(payload, "school_id", None)
    final_school_name = payload.school_name.strip() if payload.school_name else None

    if final_school_id:
        school = db.query(models.School).filter(models.School.id == final_school_id).first()
        if not school: raise HTTPException(400, "Invalid school selected")
        final_school_name = None

    if not final_school_id and not final_school_name:
        raise HTTPException(400, "School is required")

    referred_by_id = None
    referrer = None
    if payload.referral_code:
        referrer = db.query(models.User).filter(models.User.username == payload.referral_code).first()
        if referrer:
            referred_by_id = referrer.id

    trial_end = datetime.now(timezone.utc) + timedelta(days=7)

    user = models.User(
        username=username,
        display_name=payload.display_name.strip(),
        password_hash=hash_password(payload.password),
        school_id=final_school_id,
        school_name=final_school_name,
        email=payload.email.strip().lower() if payload.email else None,
        education_level=payload.education_level,
        exam_target=payload.exam_target or "jamb",
        subjects=payload.subjects,
        avatar_id=payload.avatar_id or "default",
        is_active=True,
        is_admin=False,
        referred_by_id=referred_by_id,
        trial_ends_at=trial_end
    )
    db.add(user); db.commit(); db.refresh(user)
    log_activity(db, "USER_REGISTERED", f"New student: {user.display_name} (7-Day Trial Started)")

    referral_bonus = 0
    if referrer:
        referral_bonus = 500
        referrer.coins += referral_bonus
        referrer.trial_ends_at = (referrer.trial_ends_at or datetime.now(timezone.utc)) + timedelta(days=3)
        db.commit()
        log_activity(db, "REFERRAL_REWARD", f"{referrer.display_name} earned 500 coins referring {user.display_name}")

    if user.email:
        background_tasks.add_task(email_service.send_welcome_email, user.email, user.display_name, user.username)

    token = create_access_token(user.id, user.is_admin)
    return RegisterResponse(
        user=UserOut.model_validate(user),
        token=token,
        recovery_username=user.username,
        referral_bonus=referral_bonus
    )


@app.post("/auth/login")
def login(payload: LoginRequest, request: Request, db: Session = Depends(get_db)):
    user = db.query(models.User).filter(models.User.username == payload.username).first()
    if not user or not verify_password(payload.password, user.password_hash):
        raise HTTPException(status_code=401, detail="invalid credentials")
    if not user.is_active:
        raise HTTPException(status_code=403, detail="account disabled")

    ua = request.headers.get("user-agent", "")
    device_name = "Chrome on Windows" if "Windows" in ua else ("Safari on iOS" if "iPhone" in ua else "Unknown Device")
    ip = request.client.host if request.client else None

    db.query(models.StudySession).filter(models.StudySession.user_id == user.id).update({"is_current": False})
    recent_cutoff = datetime.now(timezone.utc) - timedelta(hours=6)
    existing = db.query(models.StudySession).filter(
        models.StudySession.user_id == user.id,
        models.StudySession.device_name == device_name,
        models.StudySession.ip_address == ip,
        models.StudySession.last_active >= recent_cutoff,
    ).order_by(models.StudySession.last_active.desc()).first()
    if existing:
        existing.is_current = True
        existing.last_active = datetime.now(timezone.utc)
    else:
        db.add(models.StudySession(user_id=user.id, device_name=device_name, ip_address=ip, is_current=True))
    db.commit()

    _touch_streak(user, db)

    token = create_access_token(user.id, user.is_admin)
    if user.is_admin:
        log_activity(db, "ADMIN_LOGIN", f"Admin {user.display_name} logged in")

    return {"user": UserOut.model_validate(user), "token": token}


@app.get("/auth/me", response_model=UserOut)
def me(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    _touch_streak(user, db)
    return UserOut.model_validate(user)


@app.get("/auth/sessions")
def get_my_sessions(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    sessions = db.query(models.StudySession).filter(models.StudySession.user_id == user.id).order_by(models.StudySession.last_active.desc()).all()
    return [StudySessionOut.model_validate(s) for s in sessions]


@app.post("/auth/sessions/logout-others")
def logout_other_sessions(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    db.query(models.StudySession).filter(
        models.StudySession.user_id == user.id,
        models.StudySession.is_current == False
    ).delete(synchronize_session=False)
    db.commit()
    return {"ok": True, "message": "All other sessions logged out."}


@app.patch("/auth/me", response_model=UserOut)
def update_me(payload: ProfileUpdateRequest, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    settings = db.query(models.Settings).first()
    if settings and settings.allow_profile_edits == False:
        raise HTTPException(403, "Profile edits have been disabled by the administrator.")

    data = payload.model_dump(exclude_unset=True)
    new_level = data.get("education_level", user.education_level)

    if new_level not in EDUCATION_LEVELS:
        raise HTTPException(400, "bad education_level")

    if new_level == "primary":
        new_subjects = ["general_science"]
    else:
        new_subjects = SUBJECTS_BY_LEVEL[new_level]

    for k, v in data.items():
        if k not in ("subjects", "education_level", "exam_target"):
            setattr(user, k, v)

    old_level = user.education_level
    user.education_level = new_level
    user.subjects = new_subjects

    if new_level != old_level:
        user.exam_target = EXAM_TARGETS_BY_LEVEL[new_level][0]

    db.commit(); db.refresh(user)
    return UserOut.model_validate(user)


@app.post("/auth/forgot-password")
async def forgot_password(payload: dict, background_tasks: BackgroundTasks, db: Session = Depends(get_db)):
    email = payload.get("email", "").strip().lower()
    user = db.query(models.User).filter(models.User.email == email).first()
    if user:
        token = secrets.token_urlsafe(32)
        user.reset_token = token
        db.commit()
        background_tasks.add_task(email_service.send_password_reset_email, user.email, token)
    return {"ok": True, "message": "If an account exists, a reset link has been sent."}


@app.post("/auth/reset-password")
def reset_password(payload: dict, db: Session = Depends(get_db)):
    token = payload.get("token")
    new_password = payload.get("password")
    if not token or not new_password: raise HTTPException(400, "Missing data")
    user = db.query(models.User).filter(models.User.reset_token == token).first()
    if not user: raise HTTPException(400, "Invalid or expired token")
    user.password_hash = hash_password(new_password)
    user.reset_token = None
    db.commit()
    return {"ok": True, "message": "Password reset successfully"}


@app.post("/support")
async def contact_support(payload: dict, background_tasks: BackgroundTasks, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    message = payload.get("message", "")
    if not message: raise HTTPException(400, "Message is required")
    settings = db.query(models.Settings).first()
    support_email = settings.support_email if settings else "admin@makatron.com"
    background_tasks.add_task(email_service.send_support_email, user.display_name, user.email or "No email", message)
    log = models.EmailLog(user_id=user.id, recipient=support_email, subject="Support Request from " + user.display_name, body=message)
    db.add(log); db.commit()
    return {"ok": True, "message": "Support request sent"}


@app.get("/support/my-messages")
def get_my_support_messages(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    logs = db.query(models.EmailLog).filter(models.EmailLog.user_id == user.id).order_by(models.EmailLog.sent_at.desc()).all()
    return [{"id": l.id, "subject": l.subject, "body": l.body, "reply": l.reply, "replied_at": l.replied_at.isoformat() if l.replied_at else None, "sent_at": l.sent_at.isoformat()} for l in logs]


@app.post("/bugs/report")
def report_bug(payload: BugReportRequest, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    log = models.EmailLog(
        user_id=user.id,
        recipient="support@makaverseedtech.com",
        subject=f"Bug Report from {user.display_name}",
        body=f"Page: {payload.page_url}\nBrowser: {payload.browser_info}\n\n{payload.description}"
    )
    db.add(log); db.commit()
    log_activity(db, "BUG_REPORTED", f"{user.display_name} reported a bug: {payload.description[:80]}")
    return {"ok": True, "message": "Bug reported. Thank you!"}

# ---------------- STUDENT DASHBOARD & ARENA ----------------

@app.get("/dashboard/summary")
def dashboard_summary(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    tests = db.query(models.Test).filter(models.Test.user_id == user.id, models.Test.status == "submitted").all()
    total_tests = len(tests)
    avg_score = round(sum(t.score for t in tests) / total_tests, 1) if total_tests > 0 else 0
    weakest = None
    if total_tests > 0:
        subj_stats = {}
        for t in tests:
            if t.subject not in subj_stats: subj_stats[t.subject] = []
            subj_stats[t.subject].append(t.score)
        lowest_avg = 999
        for subj, scores in subj_stats.items():
            avg = sum(scores) / len(scores)
            if avg < lowest_avg: lowest_avg = avg; weakest = subj
    settings = db.query(models.Settings).first()
    return {"total_tests": total_tests, "average_score": avg_score, "weakest_subject": weakest, "announcement": settings.announcement if settings else None, "allow_profile_edits": settings.allow_profile_edits if settings else True, "xp": user.xp, "coins": user.coins, "total_study_time": user.total_study_time, "avatar_id": user.avatar_id, "active_title": user.active_title, "current_streak": user.current_streak or 0, "last_login_date": user.last_login_date.isoformat() if user.last_login_date else None}


@app.get("/dashboard/referral-rewards")
def referral_rewards(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    count = db.query(models.User).filter(models.User.referred_by_id == user.id).count()
    return {"referred_count": count, "coins_earned": count * 500, "premium_days_earned": count * 3}


@app.get("/my-school")
def get_my_school(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    school = None
    if user.school_id:
        school = db.query(models.School).filter(models.School.id == user.school_id).first()
    all_school_scores = db.query(models.User.school_id, models.User.school_name, func.sum(models.Test.score).label("total_points"), func.count(models.User.id.distinct()).label("student_count")).join(models.Test, models.Test.user_id == models.User.id).filter(models.Test.status == "submitted", models.Test.mode == "championship", models.User.is_admin == False).group_by(models.User.school_id, models.User.school_name).all()
    merged = {}
    for r in all_school_scores:
        key = r.school_id if r.school_id else r.school_name
        if not key: continue
        if key not in merged: merged[key] = {"name": r.school_name, "total_score": 0, "student_count": 0}
        merged[key]["total_score"] += r.total_points
        merged[key]["student_count"] += r.student_count
    sorted_schools = list(merged.values())
    sorted_schools.sort(key=lambda x: x["total_score"], reverse=True)
    my_rank = None; my_score = 0; my_students = 0
    for idx, s in enumerate(sorted_schools):
        is_mine = False
        if user.school_id and s["name"] == (school.name if school else ""): is_mine = True
        elif user.school_name and s["name"] == user.school_name: is_mine = True
        if is_mine:
            my_rank = idx + 1; my_score = s["total_score"]; my_students = s["student_count"]; break
    return {"school_name": school.name if school else (user.school_name or "Not Assigned"), "school_state": school.state if school else "—", "school_rank": my_rank, "school_score": my_score, "active_students": my_students}


@app.get("/exams/active")
def get_active_exams(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    now = datetime.now(timezone.utc)
    exams = db.query(models.Exam).filter(models.Exam.education_level == user.education_level, models.Exam.is_active == True, models.Exam.exam_type.in_(["practice", "mock"]), (models.Exam.start_date == None) | (models.Exam.start_date <= now), (models.Exam.end_date == None) | (models.Exam.end_date >= now)).order_by(models.Exam.created_at.desc()).all()
    result = []
    for e in exams:
        completed_test = None
        if e.exam_type in ("mock", "championship"):
            completed_test = db.query(models.Test).filter(models.Test.user_id == user.id, models.Test.exam_id == e.id, models.Test.status == "submitted").first()
        is_closed = bool(e.end_date and now > e.end_date)
        result.append({"id": e.id, "title": e.title, "exam_type": e.exam_type, "subject": e.subject, "duration_mins": e.duration_mins, "questions_to_serve": e.questions_to_serve, "start_date": e.start_date.isoformat() if e.start_date else None, "end_date": e.end_date.isoformat() if e.end_date else None, "completed": bool(completed_test), "completed_score": completed_test.score if completed_test else None, "completed_test_id": completed_test.id if completed_test else None, "is_closed": is_closed})
    return result


@app.get("/championship/public")
def get_public_championship(db: Session = Depends(get_db)):
    now = datetime.now(timezone.utc)
    live = db.query(models.Exam).filter(models.Exam.exam_type == "championship", models.Exam.is_active == True, models.Exam.start_date <= now, models.Exam.end_date >= now).order_by(models.Exam.end_date.asc()).first()
    if live: return {"status": "live", "title": live.title, "end_date": live.end_date.isoformat()}
    upcoming = db.query(models.Exam).filter(models.Exam.exam_type == "championship", models.Exam.is_active == True, models.Exam.start_date > now).order_by(models.Exam.start_date.asc()).first()
    if upcoming: return {"status": "upcoming", "title": upcoming.title, "start_date": upcoming.start_date.isoformat()}
    return {"status": "none"}


@app.get("/championship/current")
def get_current_championship(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    now = datetime.now(timezone.utc)
    live = db.query(models.Exam).filter(models.Exam.exam_type == "championship", models.Exam.education_level == user.education_level, models.Exam.is_active == True, models.Exam.start_date <= now, models.Exam.end_date >= now).order_by(models.Exam.end_date.asc()).first()
    if live:
        attempted = db.query(models.Test).filter(models.Test.user_id == user.id, models.Test.exam_id == live.id, models.Test.status == "submitted").first()
        return {"status": "live", "exam_id": live.id, "title": live.title, "subject": live.subject, "duration_mins": live.duration_mins, "start_date": live.start_date.isoformat(), "end_date": live.end_date.isoformat(), "attempted": bool(attempted), "attempted_score": attempted.score if attempted else None, "is_premium": bool(user.is_premium)}
    upcoming = db.query(models.Exam).filter(models.Exam.exam_type == "championship", models.Exam.education_level == user.education_level, models.Exam.is_active == True, models.Exam.start_date > now).order_by(models.Exam.start_date.asc()).first()
    if upcoming:
        return {"status": "upcoming", "exam_id": upcoming.id, "title": upcoming.title, "subject": upcoming.subject, "duration_mins": upcoming.duration_mins, "start_date": upcoming.start_date.isoformat(), "end_date": upcoming.end_date.isoformat(), "attempted": False, "attempted_score": None, "is_premium": bool(user.is_premium)}
    return {"status": "none"}


@app.get("/championship/results/{exam_id}")
def championship_results(exam_id: int, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    exam = db.query(models.Exam).filter(models.Exam.id == exam_id).first()
    if not exam: raise HTTPException(404, "Championship not found")
    now = datetime.now(timezone.utc)
    if exam.end_date and exam.end_date > now: raise HTTPException(403, "This championship has not ended yet. Results will be published after the end date.")
    top_scorers = db.query(models.Test.user_id, func.sum(models.Test.score).label("total_score"), func.sum(models.Test.time_taken_seconds).label("total_time")).filter(models.Test.exam_id == exam_id, models.Test.status == "submitted").group_by(models.Test.user_id).order_by(desc("total_score"), "total_time").limit(20).all()
    users = {u.id: u for u in db.query(models.User).filter(models.User.id.in_([s.user_id for s in top_scorers])).all()}
    return {"exam_id": exam.id, "title": exam.title, "end_date": exam.end_date.isoformat() if exam.end_date else None, "leaderboard": [{"rank": idx + 1, "user_id": s.user_id, "display_name": users[s.user_id].display_name if s.user_id in users else "Unknown", "school_name": users[s.user_id].school_name if s.user_id in users else "", "avatar_id": users[s.user_id].avatar_id if s.user_id in users else "default", "score": s.total_score} for idx, s in enumerate(top_scorers)]}


@app.get("/exam/status/{exam_id}")
def exam_status(exam_id: int, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    test = db.query(models.Test).filter(models.Test.user_id == user.id, models.Test.exam_id == exam_id, models.Test.status == "submitted").first()
    return {"completed": bool(test), "test_id": test.id if test else None, "score": test.score if test else None}


@app.get("/arena/active")
def get_active_arena(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    now = datetime.now(timezone.utc)
    exam = db.query(models.Exam).filter(models.Exam.exam_type == "championship", models.Exam.education_level == user.education_level, models.Exam.start_date <= now, models.Exam.end_date >= now, models.Exam.is_active == True).first()
    if not exam: return {"active": False}
    attempted = db.query(models.Test).filter(models.Test.user_id == user.id, models.Test.exam_id == exam.id, models.Test.status == "submitted").first()
    if attempted: return {"active": True, "attempted": True, "title": exam.title, "end_date": exam.end_date.isoformat()}
    return {"active": True, "attempted": False, "exam_id": exam.id, "title": exam.title, "duration_mins": exam.duration_mins, "end_date": exam.end_date.isoformat()}


# ---------------- TESTS ----------------

@app.post("/tests/start")
def start_test(payload: StartTestRequest, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    _abandon_stale_tests(db, user.id)

    if payload.exam_id:
        exam = db.query(models.Exam).filter(models.Exam.id == payload.exam_id, models.Exam.is_active == True).first()
        if not exam: raise HTTPException(404, "Exam not found or inactive")
        now = datetime.now(timezone.utc)
        if exam.start_date and now < exam.start_date: raise HTTPException(403, "This exam has not started yet.")
        if exam.end_date and now > exam.end_date: raise HTTPException(403, "This exam has ended.")
        if exam.access_code and payload.access_code != exam.access_code: raise HTTPException(403, "Invalid access code.")
        if exam.exam_type == "championship": require_premium(user)
        elif exam.exam_type == "mock": check_trial_access(user)
        if exam.exam_type in ("mock", "championship"):
            attempted = db.query(models.Test).filter(models.Test.user_id == user.id, models.Test.exam_id == exam.id, models.Test.status == "submitted").first()
            if attempted: raise HTTPException(403, f"You have already attempted this {exam.exam_type}. No retakes allowed.")
        pool = db.query(models.Question).filter(models.Question.exam_id == exam.id, models.Question.is_active == True).all()
        if len(pool) == 0: raise HTTPException(400, "No questions found for this exam.")
        serve_count = exam.questions_to_serve if exam.questions_to_serve > 0 else len(pool)
        if len(pool) < serve_count: raise HTTPException(400, f"Not enough questions. Exam requires {serve_count}, but only {len(pool)} exist.")
        chosen = random.sample(pool, serve_count) if exam.randomize_questions else pool[:serve_count]
        duration = exam.duration_mins * 60
        test = models.Test(user_id=user.id, exam_id=exam.id, subject=exam.subject, mode=exam.exam_type, education_level=user.education_level, question_count=len(chosen), duration_seconds=duration, status="in_progress")
    else:
        existing = db.query(models.Test).filter(models.Test.user_id == user.id, models.Test.subject == payload.subject, models.Test.status == "in_progress").first()
        if existing and existing.started_at and existing.started_at + timedelta(seconds=existing.duration_seconds) > datetime.now(timezone.utc):
            questions = db.query(models.Question).filter(models.Question.id.in_([a.question_id for a in db.query(models.TestAnswer).filter_by(test_id=existing.id).all()])).all()
            return {"test_id": existing.id, "subject": existing.subject, "mode": existing.mode, "question_count": existing.question_count, "duration_seconds": existing.duration_seconds, "started_at": existing.started_at.isoformat(), "expires_at": (existing.started_at + timedelta(seconds=existing.duration_seconds)).isoformat(), "questions": [{"id": q.id, "topic": q.topic, "difficulty": q.difficulty, "question_text": q.question_text, "image_url": q.image_url, "passage": get_passage_payload(db, q.passage_id), "options": [{"letter": "A", "text": q.option_a}, {"letter": "B", "text": q.option_b}, {"letter": "C", "text": q.option_c}, {"letter": "D", "text": q.option_d}]} for q in questions]}
        elif existing:
            existing.status = "abandoned"; existing.submitted_at = datetime.now(timezone.utc); db.commit()
        q_count = 20; duration = 1200
        filters = [models.Question.education_level == user.education_level, models.Question.subject == payload.subject, models.Question.is_active == True]
        if payload.mode == "mock": check_trial_access(user); filters.append(models.Question.exam_target == user.exam_target)
        pool = db.query(models.Question).filter(*filters).all()
        if len(pool) < q_count: raise HTTPException(400, f"Not enough general questions. Need {q_count}, found {len(pool)}.")
        chosen = random.sample(pool, q_count)
        test = models.Test(user_id=user.id, subject=payload.subject, mode=payload.mode, education_level=user.education_level, exam_target=user.exam_target if payload.mode == "mock" else None, question_count=q_count, duration_seconds=duration, status="in_progress")

    db.add(test); db.flush()
    for q in chosen: db.add(models.TestAnswer(test_id=test.id, question_id=q.id, selected_option=None, is_correct=False))
    db.commit(); db.refresh(test)
    started = test.started_at or datetime.now(timezone.utc)
    expires = started + timedelta(seconds=test.duration_seconds)

    return {"test_id": test.id, "subject": test.subject, "mode": test.mode, "question_count": test.question_count, "duration_seconds": test.duration_seconds, "started_at": started.isoformat(), "expires_at": expires.isoformat(), "questions": [{"id": q.id, "topic": q.topic, "difficulty": q.difficulty, "question_text": q.question_text, "image_url": q.image_url, "passage": get_passage_payload(db, q.passage_id), "options": [{"letter": "A", "text": q.option_a}, {"letter": "B", "text": q.option_b}, {"letter": "C", "text": q.option_c}, {"letter": "D", "text": q.option_d}]} for q in chosen]}


@app.post("/tests/{test_id}/heartbeat")
def heartbeat_test(test_id: int, payload: HeartbeatRequest, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    test = db.query(models.Test).filter(models.Test.id == test_id, models.Test.user_id == user.id).first()
    if not test: raise HTTPException(404, "test not found")
    db.query(models.StudySession).filter(models.StudySession.user_id == user.id, models.StudySession.is_current == True).update({"last_active": datetime.now(timezone.utc)})
    if test.status != "in_progress": db.commit(); return {"ok": False, "reason": "not_in_progress"}
    if test.started_at and (datetime.now(timezone.utc) - test.started_at).total_seconds() > test.duration_seconds + 30:
        test.status = "abandoned"; test.submitted_at = datetime.now(timezone.utc); db.commit()
        return {"ok": False, "reason": "time_expired"}
    if not payload.is_fullscreen:
        test.tab_switches += 1
        if test.tab_switches >= 3:
            test.status = "abandoned"; test.submitted_at = datetime.now(timezone.utc); db.commit()
            return {"ok": False, "reason": "too_many_violations"}
    db.commit()
    return {"ok": True}


@app.post("/tests/{test_id}/submit")
def submit_test(test_id: int, payload: SubmitTestRequest, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    test = db.query(models.Test).filter(models.Test.id == test_id, models.Test.user_id == user.id).first()
    if not test: raise HTTPException(404, "test not found")
    if test.status != "in_progress": raise HTTPException(400, "test already submitted")

    if test.started_at and test.duration_seconds:
        elapsed = (datetime.now(timezone.utc) - test.started_at).total_seconds()
        if elapsed > test.duration_seconds + 15:
            test.status = "submitted"; test.submitted_at = datetime.now(timezone.utc); test.score = 0; db.commit()
            raise HTTPException(403, "Time limit exceeded. Your test has been auto-submitted with a score of 0.")

    correct_count, wrong_count, skipped_count = 0, 0, 0
    answers_by_qid = {a.question_id: a for a in db.query(models.TestAnswer).filter(models.TestAnswer.test_id == test.id).all()}
    submitted = {a.question_id: a for a in payload.answers}

    for qid, ans_row in answers_by_qid.items():
        ans_obj = submitted.get(qid)
        sel = ans_obj.selected_option if ans_obj else None
        time_spent = ans_obj.time_spent_seconds if ans_obj else None
        ans_row.selected_option = sel
        q = db.query(models.Question).filter(models.Question.id == qid).first()
        if q and time_spent is not None:
            if q.time_to_answer_seconds is None:
                q.time_to_answer_seconds = time_spent
            else:
                q.time_to_answer_seconds = round((q.time_to_answer_seconds * 0.9) + (time_spent * 0.1), 1)
        if sel is None: ans_row.is_correct = False; skipped_count += 1
        elif sel == q.correct_option: ans_row.is_correct = True; correct_count += 1
        else: ans_row.is_correct = False; wrong_count += 1

    marks = 2 if test.mode == "championship" else 1
    score = correct_count * marks
    now = datetime.now(timezone.utc)
    test.score = score; test.status = "submitted"; test.tab_switches = payload.tab_switches; test.submitted_at = now; test.time_taken_seconds = int((now - test.started_at).total_seconds())
    xp_earned = 10; coins_earned = 0
    pct = (score / test.question_count) * 100 if test.question_count > 0 else 0
    if test.mode == "practice":
        if pct >= 80: coins_earned = 10
        elif pct >= 50: coins_earned = 5
        else: coins_earned = 2
    elif test.mode in ["mock", "championship"]:
        if pct >= 80: coins_earned = 30; xp_earned += 25
        elif pct >= 50: coins_earned = 15; xp_earned += 15
        else: coins_earned = 5
    user.xp += xp_earned; user.coins += coins_earned
    db.commit()
    log_activity(db, "TEST_SUBMITTED", f"{user.display_name} scored {score} in {test.subject}, earned {xp_earned} XP and {coins_earned} Coins.")
    return {"ok": True, "test_id": test.id, "score": score, "xp_earned": xp_earned, "coins_earned": coins_earned}


@app.post("/tests/{test_id}/abandon")
def abandon_test(test_id: int, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    test = db.query(models.Test).filter(models.Test.id == test_id, models.Test.user_id == user.id).first()
    if not test: raise HTTPException(404, "test not found")
    if test.status == "in_progress":
        test.status = "abandoned"; test.submitted_at = datetime.now(timezone.utc); db.commit()
    return {"ok": True}


@app.get("/tests/{test_id}/result")
def test_result(test_id: int, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    test = db.query(models.Test).filter(models.Test.id == test_id, models.Test.user_id == user.id).first()
    if not test: raise HTTPException(404, "test not found")
    if test.status == "in_progress": raise HTTPException(400, "test not submitted")
    answers = db.query(models.TestAnswer).filter(models.TestAnswer.test_id == test.id).all()
    questions = db.query(models.Question).filter(models.Question.id.in_([a.question_id for a in answers])).all()
    qmap = {q.id: q for q in questions}
    correct_count = sum(1 for a in answers if a.is_correct)
    skipped_count = sum(1 for a in answers if a.selected_option is None)
    wrong_count = len(answers) - correct_count - skipped_count
    breakdown = [{"question_id": q.id, "question_text": q.question_text, "image_url": q.image_url, "passage": get_passage_payload(db, q.passage_id), "time_to_answer_seconds": q.time_to_answer_seconds, "options": {"A": q.option_a, "B": q.option_b, "C": q.option_c, "D": q.option_d}, "selected_option": a.selected_option, "correct_option": q.correct_option, "is_correct": a.is_correct, "explanation": q.explanation} for a in answers for q in [qmap.get(a.question_id)] if q]
    return {"test_id": test.id, "subject": test.subject, "mode": test.mode, "question_count": test.question_count, "score": test.score or 0, "correct_count": correct_count, "wrong_count": wrong_count, "skipped_count": skipped_count, "time_taken_seconds": test.time_taken_seconds, "breakdown": breakdown}


@app.get("/tests-list/my")
def my_tests(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    tests = db.query(models.Test).filter(models.Test.user_id == user.id, models.Test.status != "in_progress").order_by(models.Test.submitted_at.desc()).limit(50).all()
    return [{"id": t.id, "subject": t.subject, "mode": t.mode, "score": t.score, "question_count": t.question_count, "time_taken_seconds": t.time_taken_seconds, "submitted_at": t.submitted_at.isoformat() if t.submitted_at else None} for t in tests]


@app.post("/tests/review-mistakes")
def review_mistakes(payload: Optional[ReviewMistakesRequest] = None, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    target_test = None
    if payload and payload.test_id:
        target_test = db.query(models.Test).filter(models.Test.id == payload.test_id, models.Test.user_id == user.id, models.Test.status == "submitted").first()
        if not target_test:
            raise HTTPException(404, "Test not found or not submitted.")
    if not target_test:
        target_test = db.query(models.Test).filter(models.Test.user_id == user.id, models.Test.status == "submitted").order_by(models.Test.submitted_at.desc()).first()
    if not target_test: raise HTTPException(404, "No previous test found.")
    wrong_answers = db.query(models.TestAnswer).filter(models.TestAnswer.test_id == target_test.id, models.TestAnswer.is_correct == False).all()
    if not wrong_answers: raise HTTPException(400, "You got everything right! No mistakes to review.")
    q_ids = [a.question_id for a in wrong_answers]
    questions = db.query(models.Question).filter(models.Question.id.in_(q_ids)).all()
    test = models.Test(user_id=user.id, subject=target_test.subject, mode="practice", education_level=user.education_level, question_count=len(questions), duration_seconds=len(questions)*60, status="in_progress")
    db.add(test); db.flush()
    for q in questions: db.add(models.TestAnswer(test_id=test.id, question_id=q.id, selected_option=None, is_correct=False))
    db.commit(); db.refresh(test)
    return {"test_id": test.id, "subject": test.subject, "mode": "practice", "question_count": test.question_count, "duration_seconds": test.duration_seconds, "started_at": test.started_at.isoformat(), "expires_at": (test.started_at + timedelta(seconds=test.duration_seconds)).isoformat(), "questions": [{"id": q.id, "topic": q.topic, "difficulty": q.difficulty, "question_text": q.question_text, "image_url": q.image_url, "passage": get_passage_payload(db, q.passage_id), "options": [{"letter": "A", "text": q.option_a}, {"letter": "B", "text": q.option_b}, {"letter": "C", "text": q.option_c}, {"letter": "D", "text": q.option_d}]} for q in questions]}


# ---------------- BOOKMARKS ----------------

@app.post("/bookmarks/{question_id}")
def toggle_bookmark(question_id: int, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    q = db.query(models.Question).filter(models.Question.id == question_id).first()
    if not q: raise HTTPException(404, "Question not found")
    existing = db.query(models.QuestionBookmark).filter(
        models.QuestionBookmark.user_id == user.id,
        models.QuestionBookmark.question_id == question_id
    ).first()
    if existing:
        db.delete(existing); db.commit()
        return {"bookmarked": False}
    else:
        bookmark = models.QuestionBookmark(user_id=user.id, question_id=question_id)
        db.add(bookmark); db.commit()
        return {"bookmarked": True}


@app.get("/bookmarks")
def get_my_bookmarks(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    bookmarks = db.query(models.QuestionBookmark).filter(models.QuestionBookmark.user_id == user.id).order_by(models.QuestionBookmark.created_at.desc()).all()
    question_ids = [b.question_id for b in bookmarks]
    questions = db.query(models.Question).filter(models.Question.id.in_(question_ids)).all()
    qmap = {q.id: q for q in questions}
    return [{
        "question_id": b.question_id,
        "bookmarked_at": b.created_at.isoformat(),
        "subject": qmap[b.question_id].subject if b.question_id in qmap else "unknown",
        "topic": qmap[b.question_id].topic if b.question_id in qmap else None,
        "question_text": qmap[b.question_id].question_text if b.question_id in qmap else "",
        "image_url": qmap[b.question_id].image_url if b.question_id in qmap else None,
        "options": {
            "A": qmap[b.question_id].option_a if b.question_id in qmap else "",
            "B": qmap[b.question_id].option_b if b.question_id in qmap else "",
            "C": qmap[b.question_id].option_c if b.question_id in qmap else "",
            "D": qmap[b.question_id].option_d if b.question_id in qmap else "",
        },
        "correct_option": qmap[b.question_id].correct_option if b.question_id in qmap else "",
        "explanation": qmap[b.question_id].explanation if b.question_id in qmap else None,
    } for b in bookmarks if b.question_id in qmap]


# ---------------- REPORT QUESTIONS ----------------

@app.post("/questions/{question_id}/report")
def report_question(question_id: int, payload: ReportQuestionRequest, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    q = db.query(models.Question).filter(models.Question.id == question_id).first()
    if not q: raise HTTPException(404, "Question not found")
    report = models.ReportedQuestion(question_id=question_id, user_id=user.id, reason=payload.reason, details=payload.details)
    db.add(report); db.commit()
    log_activity(db, "QUESTION_REPORTED", f"{user.display_name} reported question {question_id} for: {payload.reason}")
    return {"ok": True, "message": "Report submitted. Thank you!"}


# ---------------- SQUADS ----------------

@app.post("/squads")
def create_squad(payload: SquadCreateRequest, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    existing = db.query(models.SquadMember).filter(models.SquadMember.user_id == user.id).first()
    if existing: raise HTTPException(400, "You are already in a squad. Leave it first.")
    invite_code = secrets.token_urlsafe(6)
    squad = models.Squad(name=payload.name, invite_code=invite_code, creator_id=user.id)
    db.add(squad); db.flush()
    db.add(models.SquadMember(squad_id=squad.id, user_id=user.id))
    db.commit()
    return {"ok": True, "squad": {"id": squad.id, "name": squad.name, "invite_code": invite_code}}


@app.post("/squads/join")
def join_squad(payload: SquadJoinRequest, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    existing = db.query(models.SquadMember).filter(models.SquadMember.user_id == user.id).first()
    if existing: raise HTTPException(400, "You are already in a squad.")
    squad = db.query(models.Squad).filter(models.Squad.invite_code == payload.invite_code).first()
    if not squad: raise HTTPException(404, "Invalid invite code.")
    member_count = db.query(models.SquadMember).filter(models.SquadMember.squad_id == squad.id).count()
    if member_count >= 10: raise HTTPException(400, "This squad is full (max 10 members).")
    db.add(models.SquadMember(squad_id=squad.id, user_id=user.id))
    db.commit()
    return {"ok": True, "squad": {"id": squad.id, "name": squad.name}}


@app.get("/squads/my")
def get_my_squad(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    membership = db.query(models.SquadMember).filter(models.SquadMember.user_id == user.id).first()
    if not membership: return {"in_squad": False}
    squad = db.query(models.Squad).filter(models.Squad.id == membership.squad_id).first()
    members = db.query(models.SquadMember).filter(models.SquadMember.squad_id == squad.id).all()
    leaderboard = []
    for m in members:
        member_user = db.query(models.User).filter(models.User.id == m.user_id).first()
        total_score = db.query(func.sum(models.Test.score)).filter(models.Test.user_id == m.user_id, models.Test.mode == "championship", models.Test.status == "submitted").scalar() or 0
        leaderboard.append({"user_id": member_user.id, "display_name": member_user.display_name, "avatar_id": member_user.avatar_id, "total_score": total_score})
    leaderboard.sort(key=lambda x: x["total_score"], reverse=True)
    return {"in_squad": True, "squad": {"id": squad.id, "name": squad.name, "invite_code": squad.invite_code, "creator_id": squad.creator_id}, "members": leaderboard}


@app.post("/squads/leave")
def leave_squad(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    membership = db.query(models.SquadMember).filter(models.SquadMember.user_id == user.id).first()
    if not membership: raise HTTPException(404, "You are not in a squad.")
    db.delete(membership); db.commit()
    return {"ok": True, "message": "You have left the squad."}


# ---------------- LEADERBOARD ----------------

@app.get("/leaderboard/schools")
def lb_schools(db: Session = Depends(get_db)):
    results = db.query(
        models.User.school_id, models.User.school_name,
        func.sum(models.Test.score).label("total_points"),
        func.count(models.User.id.distinct()).label("student_count")
    ).join(models.Test, models.Test.user_id == models.User.id).filter(
        models.Test.status == "submitted", models.Test.mode == "championship",
        models.User.is_admin == False
    ).group_by(models.User.school_id, models.User.school_name).all()

    schools_map = {s.id: s.name for s in db.query(models.School).all()}
    merged = {}
    for r in results:
        key = r.school_id if r.school_id else r.school_name
        if not key: continue
        name = schools_map.get(r.school_id) if r.school_id else r.school_name
        if key not in merged:
            merged[key] = {"name": name, "total_score": 0, "student_count": 0}
        merged[key]["total_score"] += r.total_points
        merged[key]["student_count"] += r.student_count
    final_list = list(merged.values())
    final_list.sort(key=lambda x: x["total_score"], reverse=True)
    return {"rows": final_list[:20]}


@app.get("/leaderboard/state/{state_name}")
def lb_state(state_name: str, db: Session = Depends(get_db)):
    state_schools = db.query(models.School).filter(func.lower(models.School.state) == state_name.lower()).all()
    school_ids = [s.id for s in state_schools]
    if not school_ids: return {"rows": [], "state": state_name}
    results = db.query(
        models.User.school_id, models.User.school_name,
        func.sum(models.Test.score).label("total_points"),
        func.count(models.User.id.distinct()).label("student_count")
    ).join(models.Test, models.Test.user_id == models.User.id).filter(
        models.Test.status == "submitted", models.Test.mode == "championship",
        models.User.school_id.in_(school_ids)
    ).group_by(models.User.school_id, models.User.school_name).all()
    schools_map = {s.id: s.name for s in state_schools}
    merged = {}
    for r in results:
        key = r.school_id
        if not key: continue
        name = schools_map.get(r.school_id)
        if key not in merged: merged[key] = {"name": name, "total_score": 0, "student_count": 0}
        merged[key]["total_score"] += r.total_points
        merged[key]["student_count"] += r.student_count
    final_list = list(merged.values())
    final_list.sort(key=lambda x: x["total_score"], reverse=True)
    return {"rows": final_list[:20], "state": state_name}


@app.get("/leaderboard/subject/{subject}")
def lb_subject(subject: str, education_level: str = Query(...), db: Session = Depends(get_db)):
    valid_users = db.query(
        models.Test.user_id,
        func.sum(models.Test.score).label("total_points"),
        func.sum(models.Test.time_taken_seconds).label("total_time")
    ).join(models.User, models.User.id == models.Test.user_id).filter(
        models.Test.status == "submitted", models.Test.mode == "championship",
        models.Test.subject == subject, models.Test.education_level == education_level,
        models.User.is_admin == False
    ).group_by(models.Test.user_id).all()
    results = []
    for u in valid_users:
        results.append({"user_id": u.user_id, "total_score": u.total_points, "total_time": u.total_time or 999999})
    results.sort(key=lambda x: (-x["total_score"], x["total_time"]))
    top_users = db.query(models.User).filter(models.User.id.in_([r["user_id"] for r in results[:15]])).all()
    umap = {u.id: u for u in top_users}
    out = []
    for r in results[:15]:
        u = umap[r["user_id"]]
        out.append({
            "user_id": u.id, "username": u.username, "display_name": u.display_name,
            "school_name": u.school_name, "avatar_id": u.avatar_id, "total_score": r["total_score"]
        })
    return {"rows": out}


@app.get("/leaderboard/homepage")
def lb_homepage(db: Session = Depends(get_db)):
    result = {}
    schools_map = {s.id: s for s in db.query(models.School).all()}
    for target in ALL_EXAM_TARGETS:
        valid_users = db.query(
            models.Test.user_id,
            func.sum(models.Test.score).label("total_points"),
            func.sum(models.Test.time_taken_seconds).label("total_time")
        ).join(models.User, models.User.id == models.Test.user_id).filter(
            models.Test.status == "submitted", models.Test.mode == "championship",
            models.Test.exam_target == target,
            models.User.is_admin == False
        ).group_by(models.Test.user_id).all()
        results = []
        for u in valid_users:
            results.append({"user_id": u.user_id, "total_score": u.total_points, "total_time": u.total_time or 999999})
        results.sort(key=lambda x: (-x["total_score"], x["total_time"]))
        top_users = db.query(models.User).filter(models.User.id.in_([r["user_id"] for r in results[:3]])).all()
        umap = {u.id: u for u in top_users}
        out = []
        for r in results[:3]:
            u = umap[r["user_id"]]
            state = None
            if u.school_id and u.school_id in schools_map:
                state = schools_map[u.school_id].state
            out.append({
                "user_id": u.id, "display_name": u.display_name,
                "school_name": u.school_name, "school_state": state,
                "avatar_id": u.avatar_id,
                "total_score": r["total_score"]
            })
        result[target] = out
    return {"boards": result}


@app.get("/leaderboard/championship")
def lb_championship(db: Session = Depends(get_db)):
    now = datetime.now(timezone.utc)
    # Get all exams that have officially ended
    past_exams = db.query(models.Exam).filter(
        models.Exam.exam_type == "championship", 
        models.Exam.end_date < now
    ).order_by(models.Exam.end_date.desc()).all()
    
    hall_of_fame = []
    
    for exam in past_exams:
        # Fetch the top 3 scorers for this exam (tie-breaker: fastest time)
        top_scorers = db.query(
            models.Test.user_id, 
            func.sum(models.Test.score).label("total_score"),
            func.sum(models.Test.time_taken_seconds).label("total_time")
        ).filter(
            models.Test.exam_id == exam.id, 
            models.Test.status == "submitted"
        ).group_by(models.Test.user_id).order_by(
            desc("total_score"), "total_time"
        ).limit(3).all()
        
        if not top_scorers:
            continue
            
        winners_list = []
        for idx, s in enumerate(top_scorers):
            u = db.query(models.User).filter(models.User.id == s.user_id).first()
            if u:
                winners_list.append({
                    "rank": idx + 1,
                    "display_name": u.display_name,
                    "school_name": u.school_name or "Independent",
                    "avatar_id": u.avatar_id,
                    "score": s.total_score
                })
        
        if winners_list:
            hall_of_fame.append({
                "exam_id": exam.id,
                "exam_title": exam.title,
                "end_date": exam.end_date.isoformat() if exam.end_date else None,
                "winners": winners_list
            })
            
    return hall_of_fame


@app.get("/leaderboard/my-rank")
def my_rank(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    scores = db.query(
        models.User.id,
        func.coalesce(func.sum(models.Test.score), 0).label("total_score")
    ).join(models.Test, models.Test.user_id == models.User.id).filter(
        models.Test.status == "submitted", models.Test.mode == "championship",
        models.User.is_admin == False,
    ).group_by(models.User.id).all()
    valid_scores = sorted(scores, key=lambda x: x.total_score, reverse=True)
    my_position = -1
    my_score = 0
    for idx, s in enumerate(valid_scores):
        if s.id == user.id:
            my_position = idx + 1
            my_score = s.total_score
            break
    return {"rank": my_position if my_position > 0 else None, "score": my_score, "total_ranked_users": len(valid_scores)}


# ---------------- DYNAMIC TICKER ----------------

@app.get("/ticker")
def get_public_ticker(db: Session = Depends(get_db)):
    headlines = db.query(models.TickerNews).filter(models.TickerNews.is_active == True).order_by(models.TickerNews.created_at.desc()).all()
    if not headlines:
        return [
            {"id": 0, "headline": "Welcome to MAKAVERSE! The ultimate CBT preparation platform."},
            {"id": 0, "headline": "New mock exams are added weekly. Keep practicing!"},
            {"id": 0, "headline": "Compete in the Championship Arena to earn your spot in the Hall of Fame."},
        ]
    return [{"id": h.id, "headline": h.headline} for h in headlines]


@app.post("/admin/ticker")
def admin_create_ticker(payload: dict, db: Session = Depends(get_db), admin: models.User = Depends(require_super_admin)):
    headline = payload.get("headline", "").strip()
    if not headline: raise HTTPException(400, "Headline cannot be empty")
    news = models.TickerNews(headline=headline, is_active=True)
    db.add(news); db.commit(); db.refresh(news)
    log_activity(db, "TICKER_CREATED", f"Admin created ticker: {headline}")
    return {"ok": True, "id": news.id}


@app.get("/admin/ticker")
def admin_get_ticker(db: Session = Depends(get_db), admin: models.User = Depends(require_super_admin)):
    headlines = db.query(models.TickerNews).order_by(models.TickerNews.created_at.desc()).all()
    return [{"id": h.id, "headline": h.headline, "is_active": h.is_active, "created_at": h.created_at.isoformat()} for h in headlines]


@app.put("/admin/ticker/{ticker_id}/toggle")
def admin_toggle_ticker(ticker_id: int, db: Session = Depends(get_db), admin: models.User = Depends(require_super_admin)):
    news = db.query(models.TickerNews).filter(models.TickerNews.id == ticker_id).first()
    if not news: raise HTTPException(404, "Ticker not found")
    news.is_active = not news.is_active
    db.commit()
    return {"ok": True, "is_active": news.is_active}


@app.delete("/admin/ticker/{ticker_id}")
def admin_delete_ticker(ticker_id: int, db: Session = Depends(get_db), admin: models.User = Depends(require_super_admin)):
    news = db.query(models.TickerNews).filter(models.TickerNews.id == ticker_id).first()
    if not news: raise HTTPException(404, "Ticker not found")
    db.delete(news); db.commit()
    log_activity(db, "TICKER_DELETED", f"Admin deleted ticker ID: {ticker_id}")
    return {"ok": True}


# ---------------- ADMIN ----------------

@app.get("/admin/analytics")
def admin_analytics(db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    scoped = is_school_scoped(admin)

    q_base = db.query(models.User).filter(models.User.is_admin == False)
    if scoped:
        q_base = q_base.filter(models.User.school_id == admin.school_id)
    total_students = q_base.count()

    t_base = db.query(models.Test).filter(models.Test.status == "submitted")
    if scoped:
        t_base = t_base.join(models.User, models.User.id == models.Test.user_id).filter(models.User.school_id == admin.school_id)
    total_tests = t_base.count()
    avg_score_result = t_base.with_entities(func.avg(models.Test.score)).scalar()
    avg_score = round(avg_score_result, 1) if avg_score_result else 0

    today = datetime.now(timezone.utc).date()
    chart_data = []
    for i in range(6, -1, -1):
        day = today - timedelta(days=i)
        next_day = day + timedelta(days=1)
        cq = t_base.filter(
            models.Test.submitted_at >= datetime.combine(day, datetime.min.time()).replace(tzinfo=timezone.utc),
            models.Test.submitted_at < datetime.combine(next_day, datetime.min.time()).replace(tzinfo=timezone.utc)
        )
        chart_data.append({"date": day.strftime("%a"), "tests": cq.count()})

    # FIXED: Replaced func.cast with case() which is universally supported
    wt_query = db.query(
        models.Question.topic,
        func.count(models.TestAnswer.id).label("total"),
        func.sum(case((models.TestAnswer.is_correct == True, 1), else_=0)).label("correct"),
    ).join(
        models.TestAnswer, models.TestAnswer.question_id == models.Question.id
    ).join(
        models.Test, models.Test.id == models.TestAnswer.test_id
    ).filter(
        models.Test.status == "submitted",
        models.Question.topic != None,
        models.Question.topic != "",
    )
    if scoped:
        wt_query = wt_query.join(models.User, models.User.id == models.Test.user_id).filter(models.User.school_id == admin.school_id)
    wt_rows = wt_query.group_by(models.Question.topic).all()

    weakest_topics = []
    for topic, total, correct in wt_rows:
        if total and total > 0:
            correct_val = int(correct or 0)
            pct = round((correct_val / total) * 100, 1)
            if pct < 60:
                weakest_topics.append({"topic": topic, "avg_score": pct, "attempts": total})
    weakest_topics.sort(key=lambda x: x["avg_score"])
    weakest_topics = weakest_topics[:10]

    return {
        "total_students": total_students,
        "total_tests": total_tests,
        "avg_score": avg_score,
        "chart_data": chart_data,
        "weakest_topics": weakest_topics,
    }


@app.get("/admin/analytics/question-difficulty")
def admin_question_difficulty(db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    too_easy = db.query(models.Question).filter(models.Question.time_to_answer_seconds != None, models.Question.time_to_answer_seconds < 5).order_by(models.Question.time_to_answer_seconds.asc()).limit(20).all()
    too_hard = db.query(models.Question).filter(models.Question.time_to_answer_seconds != None, models.Question.time_to_answer_seconds > 180).order_by(models.Question.time_to_answer_seconds.desc()).limit(20).all()
    return {
        "too_easy": [{"id": q.id, "text": q.question_text[:80], "time": q.time_to_answer_seconds} for q in too_easy],
        "too_hard": [{"id": q.id, "text": q.question_text[:80], "time": q.time_to_answer_seconds} for q in too_hard],
    }


@app.get("/admin/settings")
def get_settings(db: Session = Depends(get_db), admin: models.User = Depends(require_super_admin)):
    s = db.query(models.Settings).first()
    if not s:
        s = models.Settings(); db.add(s); db.commit(); db.refresh(s)
    return s


@app.put("/admin/settings")
def update_settings(payload: SettingsUpdate, db: Session = Depends(get_db), admin: models.User = Depends(require_super_admin)):
    s = db.query(models.Settings).first()
    if not s: s = models.Settings()
    s.platform_name = payload.platform_name
    s.support_email = payload.support_email
    s.registration_open = payload.registration_open
    s.maintenance_mode = payload.maintenance_mode
    s.announcement = payload.announcement
    s.allow_profile_edits = payload.allow_profile_edits
    db.commit(); db.refresh(s)
    log_activity(db, "SETTINGS_UPDATED", "Admin updated global settings")
    return s


@app.get("/admin/logs")
def get_logs(db: Session = Depends(get_db), admin: models.User = Depends(require_super_admin)):
    logs = db.query(models.ActivityLog).order_by(models.ActivityLog.created_at.desc()).limit(10).all()
    return [{"action": l.action_type, "msg": l.description, "time": l.created_at.isoformat()} for l in logs]

@app.get("/admin/results")
def admin_get_results(
    db: Session = Depends(get_db), 
    admin: models.User = Depends(require_admin),
    mode: Optional[str] = None,
    education_level: Optional[str] = None,
    state: Optional[str] = None,
    days: Optional[int] = None
):
    # Base Query: Join Test, User, and School
    query = db.query(
        models.Test.id,
        models.Test.subject,
        models.Test.mode,
        models.Test.score,
        models.Test.question_count,
        models.Test.time_taken_seconds,
        models.Test.submitted_at,
        models.User.display_name,
        models.User.school_name,
        models.School.state
    ).join(
        models.User, models.Test.user_id == models.User.id
    ).outerjoin(
        models.School, models.User.school_id == models.School.id
    ).filter(
        models.Test.status == "submitted"
    )

    # Apply Filters
    if mode:
        query = query.filter(models.Test.mode == mode)
    if education_level:
        query = query.filter(models.Test.education_level == education_level)
    if state:
        query = query.filter(models.School.state == state)
    if days:
        cutoff = datetime.now(timezone.utc) - timedelta(days=days)
        query = query.filter(models.Test.submitted_at >= cutoff)

    # Execute and order by newest first
    results = query.order_by(models.Test.submitted_at.desc()).limit(500).all()

    out = []
    for r in results:
        # Calculate percentage safely
        pct = round((r.score / r.question_count) * 100, 1) if r.question_count > 0 else 0
        out.append({
            "id": r.id,
            "student_name": r.display_name,
            "school_name": r.school_name or "Independent",
            "state": r.state or "—",
            "subject": r.subject,
            "mode": r.mode,
            "score": r.score,
            "total_questions": r.question_count,
            "percentage": pct,
            "time_taken_seconds": r.time_taken_seconds,
            "date": r.submitted_at.isoformat() if r.submitted_at else None
        })
    return out


@app.post("/admin/notifications/send")
def admin_send_notification(payload: dict, db: Session = Depends(get_db), admin: models.User = Depends(require_super_admin)):
    title = payload.get("title", "Announcement")
    message = payload.get("message", "")
    target = payload.get("target", "all")
    if not message: raise HTTPException(400, "Message required")
    q = db.query(models.User).filter(models.User.email != None, models.User.is_admin == False)
    if target == "premium": q = q.filter(models.User.is_premium == True)
    elif target == "free": q = q.filter(models.User.is_premium == False)
    users = q.all()
    for u in users:
        try: email_service.send_support_email(u.display_name, u.email, f"📢 {title}\n\n{message}")
        except: pass
    log_activity(db, "NOTIFICATION_SENT", f"Admin sent '{title}' to {len(users)} users")
    return {"ok": True, "sent": len(users)}


@app.get("/admin/exams")
def admin_list_exams(db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    return db.query(models.Exam).order_by(models.Exam.created_at.desc()).all()


@app.post("/admin/exams")
def admin_create_exam(payload: dict, db: Session = Depends(get_db), admin: models.User = Depends(require_super_admin)):
    start = None
    end = None
    if payload.get("start_date"): start = datetime.fromisoformat(payload.get("start_date").replace("Z", "+00:00"))
    if payload.get("end_date"): end = datetime.fromisoformat(payload.get("end_date").replace("Z", "+00:00"))
    exam = models.Exam(
        title=payload.get("title"), exam_type=payload.get("exam_type", "mock"),
        education_level=payload.get("education_level"), subject=payload.get("subject", "mixed"),
        duration_mins=payload.get("duration_mins", 60), pass_percent=payload.get("pass_percent", 40),
        total_marks=payload.get("total_marks", 100),
        start_date=start, end_date=end,
        is_active=payload.get("is_active", True),
        add_to_bank=payload.get("add_to_bank", False),
        questions_to_serve=payload.get("questions_to_serve", 0),
        randomize_questions=payload.get("randomize_questions", True),
        access_code=payload.get("access_code"),
        instructions=payload.get("instructions")
    )
    db.add(exam); db.commit()
    log_activity(db, "EXAM_CREATED", f"Admin created {exam.exam_type} exam: {exam.title}")
    return {"ok": True, "id": exam.id}


@app.put("/admin/exams/{exam_id}")
def admin_update_exam(exam_id: int, payload: dict, db: Session = Depends(get_db), admin: models.User = Depends(require_super_admin)):
    exam = db.query(models.Exam).filter(models.Exam.id == exam_id).first()
    if not exam: raise HTTPException(404, "Exam not found")
    if "title" in payload: exam.title = payload["title"]
    if "exam_type" in payload: exam.exam_type = payload["exam_type"]
    if "education_level" in payload: exam.education_level = payload["education_level"]
    if "subject" in payload: exam.subject = payload["subject"]
    if "duration_mins" in payload: exam.duration_mins = payload["duration_mins"]
    if "pass_percent" in payload: exam.pass_percent = payload["pass_percent"]
    if "total_marks" in payload: exam.total_marks = payload["total_marks"]
    if "is_active" in payload: exam.is_active = payload["is_active"]
    if "questions_to_serve" in payload: exam.questions_to_serve = payload["questions_to_serve"]
    if "randomize_questions" in payload: exam.randomize_questions = payload["randomize_questions"]
    if "access_code" in payload: exam.access_code = payload["access_code"]
    if "instructions" in payload: exam.instructions = payload["instructions"]
    if "start_date" in payload:
        exam.start_date = datetime.fromisoformat(payload["start_date"].replace("Z", "+00:00")) if payload["start_date"] else None
    if "end_date" in payload:
        exam.end_date = datetime.fromisoformat(payload["end_date"].replace("Z", "+00:00")) if payload["end_date"] else None
    db.commit()
    log_activity(db, "EXAM_UPDATED", f"Admin updated exam: {exam.title}")
    return {"ok": True}


@app.delete("/admin/exams/{exam_id}")
def admin_delete_exam(exam_id: int, db: Session = Depends(get_db), admin: models.User = Depends(require_super_admin)):
    exam = db.query(models.Exam).filter(models.Exam.id == exam_id).first()
    if not exam: raise HTTPException(404, "Exam not found")
    exam_title = exam.title
    test_ids = [t.id for t in db.query(models.Test.id).filter(models.Test.exam_id == exam_id).all()]
    if test_ids:
        db.query(models.TestAnswer).filter(models.TestAnswer.test_id.in_(test_ids)).delete(synchronize_session=False)
    db.query(models.Test).filter(models.Test.exam_id == exam_id).delete(synchronize_session=False)
    db.query(models.HallOfFame).filter(models.HallOfFame.exam_id == exam_id).delete(synchronize_session=False)
    db.query(models.Question).filter(models.Question.exam_id == exam_id).delete(synchronize_session=False)
    db.delete(exam); db.commit()
    log_activity(db, "EXAM_DELETED", f"Admin deleted exam: {exam_title}")
    return {"ok": True}


@app.put("/admin/exams/{exam_id}/toggle")
def toggle_exam(exam_id: int, db: Session = Depends(get_db), admin: models.User = Depends(require_super_admin)):
    exam = db.query(models.Exam).filter(models.Exam.id == exam_id).first()
    if not exam: raise HTTPException(404, "Exam not found")
    exam.is_active = not exam.is_active
    db.commit()
    return {"ok": True, "is_active": exam.is_active}


@app.get("/admin/exams/{exam_id}/questions")
def admin_get_exam_questions(exam_id: int, db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    qs = db.query(models.Question).filter(models.Question.exam_id == exam_id).all()
    return [QuestionOut.model_validate(q).model_dump() for q in qs]


@app.put("/admin/questions/{question_id}")
def admin_edit_question(question_id: int, payload: QuestionUpsert, db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    q = db.query(models.Question).filter(models.Question.id == question_id).first()
    if not q: raise HTTPException(404, "Question not found")
    q.question_text = payload.question_text
    q.option_a = payload.option_a
    q.option_b = payload.option_b
    q.option_c = payload.option_c
    q.option_d = payload.option_d
    q.correct_option = payload.correct_option
    q.explanation = payload.explanation
    q.image_url = payload.image_url
    q.passage_id = payload.passage_id
    db.commit()
    log_activity(db, "QUESTION_EDITED", f"Admin edited question ID: {question_id}")
    return {"ok": True}


@app.post("/admin/questions/upload")
async def admin_upload_csv(
    file: UploadFile = File(...),
    default_level: Optional[str] = Query(None),
    default_subject: Optional[str] = Query(None),
    default_target: Optional[str] = Query(None),
    default_exam_id: Optional[int] = Query(None),
    default_is_active: bool = Query(True),
    add_to_bank: bool = Query(False),
    db: Session = Depends(get_db), admin: models.User = Depends(require_admin)
):
    filename = file.filename.lower()
    if not (filename.endswith(".csv") or filename.endswith(".xlsx")):
        raise HTTPException(400, "File must be .csv or .xlsx")
    contents = await file.read()
    rows_data = []
    if filename.endswith(".csv"):
        try: text = contents.decode("utf-8-sig")
        except: text = contents.decode("latin-1")
        reader = csv.DictReader(io.StringIO(text))
        for r in reader: rows_data.append(r)
    else:
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
        sheet = wb.active
        headers = [str(cell.value).strip().lower() if cell.value else "" for cell in sheet[1]]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not any(row): continue
            row_dict = {headers[i]: str(row[i]).strip() if row[i] is not None else "" for i in range(len(headers)) if i < len(row)}
            rows_data.append(row_dict)

    added = 0
    errors = []
    for row_num, row in enumerate(rows_data, start=2):
        try:
            def get_val(key, default=None):
                if key in row and row[key].strip(): return row[key].strip()
                return default
            req = ["question_text", "option_a", "option_b", "option_c", "option_d", "correct_option"]
            for r in req:
                if r not in row or not row[r].strip(): raise ValueError(f"Missing {r}")
            level = (get_val("education_level") or default_level or "").lower()
            subj = (get_val("subject") or default_subject or "").lower()
            target = (get_val("exam_target") or default_target or "").lower()
            if not level: raise ValueError("Missing education_level")
            if not subj: raise ValueError("Missing subject")
            if not target or target == "none": target = None
            correct = row["correct_option"].strip().upper()
            if correct not in ["A", "B", "C", "D"]: raise ValueError("Correct option must be A/B/C/D")
            q = models.Question(
                exam_id=default_exam_id,
                education_level=level, subject=subj, exam_target=target,
                topic=row.get("topic", "").strip() or None,
                difficulty=row.get("difficulty", "medium").strip().lower() or "medium",
                question_text=row["question_text"].strip(),
                option_a=row["option_a"].strip(), option_b=row["option_b"].strip(),
                option_c=row["option_c"].strip(), option_d=row["option_d"].strip(),
                correct_option=correct,
                explanation=row.get("explanation", "").strip() or None,
                image_url=row.get("image_url", "").strip() or None,
                is_active=default_is_active
            )
            db.add(q); added += 1
        except Exception as e:
            errors.append(f"Row {row_num}: {str(e)}")
    db.commit()
    log_activity(db, "QUESTIONS_UPLOADED", f"Admin uploaded {added} questions")
    return {"added": added, "errors": errors}


@app.get("/admin/questions")
def admin_list_questions(education_level: Optional[str] = None, subject: Optional[str] = None, limit: int = 100, db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    q = db.query(models.Question)
    if education_level: q = q.filter(models.Question.education_level == education_level)
    if subject: q = q.filter(models.Question.subject == subject)
    rows = q.order_by(models.Question.id.desc()).limit(limit).all()
    return {"total": len(rows), "items": [QuestionOut.model_validate(r).model_dump() for r in rows]}


# ---------------- ADMIN: PASSAGES ----------------

@app.get("/admin/passages")
def admin_list_passages(db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    rows = db.query(models.Passage).order_by(models.Passage.created_at.desc()).all()
    return [PassageOut.model_validate(p).model_dump() for p in rows]


@app.post("/admin/passages")
def admin_create_passage(payload: PassageUpsert, db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    p = models.Passage(title=payload.title, content=payload.content, subject=payload.subject, education_level=payload.education_level, exam_target=payload.exam_target)
    db.add(p); db.commit(); db.refresh(p)
    return {"ok": True, "id": p.id}


@app.delete("/admin/passages/{passage_id}")
def admin_delete_passage(passage_id: int, db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    p = db.query(models.Passage).filter(models.Passage.id == passage_id).first()
    if not p: raise HTTPException(404, "Passage not found")
    db.query(models.Question).filter(models.Question.passage_id == passage_id).update({"passage_id": None})
    db.delete(p); db.commit()
    return {"ok": True}


# ---------------- ADMIN: STAGING & REPORTS ----------------

@app.get("/admin/reports")
def admin_get_reports(db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    reports = db.query(models.ReportedQuestion).filter(models.ReportedQuestion.status == "pending").order_by(models.ReportedQuestion.created_at.desc()).all()
    return [{"id": r.id, "question_id": r.question_id, "reason": r.reason, "details": r.details, "created_at": r.created_at.isoformat()} for r in reports]


@app.post("/admin/reports/{report_id}/resolve")
def admin_resolve_report(report_id: int, db: Session = Depends(get_db), admin: models.User = Depends(require_super_admin)):
    report = db.query(models.ReportedQuestion).filter(models.ReportedQuestion.id == report_id).first()
    if not report: raise HTTPException(404, "Report not found")
    report.status = "resolved"
    db.commit()
    return {"ok": True}


@app.get("/admin/staging")
def admin_get_staging(db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    staging = db.query(models.StagingQuestion).filter(models.StagingQuestion.status == "pending").order_by(models.StagingQuestion.created_at.desc()).all()
    return [StagingQuestionOut.model_validate(q).model_dump() for q in staging]


@app.post("/admin/staging/{staging_id}/approve")
def admin_approve_staging(staging_id: int, db: Session = Depends(get_db), admin: models.User = Depends(require_super_admin)):
    sq = db.query(models.StagingQuestion).filter(models.StagingQuestion.id == staging_id).first()
    if not sq: raise HTTPException(404, "Staging question not found")
    q = models.Question(education_level=sq.education_level, subject=sq.subject, exam_target=sq.exam_target, topic=sq.topic, difficulty=sq.difficulty, question_text=sq.question_text, option_a=sq.option_a, option_b=sq.option_b, option_c=sq.option_c, option_d=sq.option_d, correct_option=sq.correct_option, explanation=sq.explanation, image_url=sq.image_url, is_active=True)
    db.add(q); sq.status = "approved"; db.commit()
    return {"ok": True}


@app.post("/admin/staging/{staging_id}/reject")
def admin_reject_staging(staging_id: int, db: Session = Depends(get_db), admin: models.User = Depends(require_super_admin)):
    sq = db.query(models.StagingQuestion).filter(models.StagingQuestion.id == staging_id).first()
    if not sq: raise HTTPException(404, "Staging question not found")
    sq.status = "rejected"; db.commit()
    return {"ok": True}


# ---------------- ADMIN: LIBRARY ----------------

@app.get("/library")
def get_study_materials(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    check_trial_access(user)
    materials = db.query(StudyMaterial).filter(StudyMaterial.is_active == True).order_by(StudyMaterial.created_at.desc()).all()
    return [StudyMaterialOut.model_validate(m).model_dump() for m in materials]


@app.post("/admin/library")
def admin_add_study_material(payload: StudyMaterialUpsert, db: Session = Depends(get_db), admin: models.User = Depends(require_super_admin)):
    material = models.StudyMaterial(
        title=payload.title, subject=payload.subject, education_level=payload.education_level,
        exam_target=payload.exam_target, drive_file_id=payload.drive_file_id,
        description=payload.description, is_active=payload.is_active
    )
    db.add(material); db.commit(); db.refresh(material)
    log_activity(db, "LIBRARY_ADDED", f"Admin added PDF: {material.title}")
    return {"ok": True, "id": material.id}


@app.get("/admin/library")
def admin_list_study_materials(db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    materials = db.query(StudyMaterial).order_by(StudyMaterial.created_at.desc()).all()
    return [StudyMaterialOut.model_validate(m).model_dump() for m in materials]


@app.delete("/admin/library/{material_id}")
def admin_delete_study_material(material_id: int, db: Session = Depends(get_db), admin: models.User = Depends(require_super_admin)):
    material = db.query(models.StudyMaterial).filter(models.StudyMaterial.id == material_id).first()
    if not material: raise HTTPException(404, "Material not found")
    db.delete(material); db.commit()
    log_activity(db, "LIBRARY_DELETED", f"Admin deleted PDF ID: {material_id}")
    return {"ok": True}


# ---------------- ADMIN: STUDENTS ----------------

@app.get("/admin/students")
def admin_list_students(db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    week_start = week_start_lagos()
    q = db.query(models.User, models.School.name).outerjoin(models.School, models.User.school_id == models.School.id).filter(models.User.is_admin == False)
    if is_school_scoped(admin):
        q = q.filter(models.User.school_id == admin.school_id)
    students = q.order_by(models.User.created_at.desc()).all()

    result = []
    for s, school_name in students:
        display_school = school_name if school_name else s.school_name
        tests_all = db.query(models.Test).filter(models.Test.user_id == s.id, models.Test.status != "in_progress").all()
        tests_week = [t for t in tests_all if t.submitted_at and t.submitted_at >= week_start]
        def pct(ts):
            if not ts: return None
            earned = sum(t.score or 0 for t in ts)
            max_possible = sum(t.question_count for t in ts)
            if max_possible == 0: return None
            return round(earned / max_possible * 100, 1)
        result.append({
            "id": s.id, "username": s.username, "display_name": s.display_name,
            "school_id": s.school_id, "school_name": display_school,
            "education_level": s.education_level,
            "exam_target": s.exam_target, "subjects": s.subjects,
            "total_tests": len(tests_all), "tests_this_week": len(tests_week),
            "overall_pct": pct(tests_all), "weekly_pct": pct(tests_week),
            "is_active": s.is_active, "email": s.email,
            "xp": s.xp, "coins": s.coins, "is_premium": s.is_premium,
            "role": s.role,
            "created_at": s.created_at.isoformat() if s.created_at else None,
        })
    return result


@app.patch("/admin/students/{user_id}/stats")
def admin_update_student_stats(user_id: int, payload: AdminStudentUpdate, db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    target = _scoped_user_or_403(db, admin, user_id)
    if payload.xp is not None: target.xp = payload.xp
    if payload.coins is not None: target.coins = payload.coins
    if payload.is_premium is not None: target.is_premium = payload.is_premium
    if payload.is_active is not None: target.is_active = payload.is_active
    if payload.role is not None:
        if payload.role in ("school_admin", "super_admin") and admin.role != "super_admin":
            raise HTTPException(403, "Only a super admin can assign this role.")
        target.role = payload.role
    if payload.school_id is not None:
        if admin.role != "super_admin":
            raise HTTPException(403, "Only a super admin can change a user's school.")
        target.school_id = payload.school_id
    db.commit()
    log_activity(db, "STUDENT_STATS_UPDATED", f"Admin updated stats for {target.display_name}")
    return {"ok": True}


@app.post("/admin/users/{user_id}/role")
def admin_update_user_role(user_id: int, payload: RoleUpdateRequest, db: Session = Depends(get_db), admin: models.User = Depends(require_super_admin)):
    target = db.query(models.User).filter(models.User.id == user_id).first()
    if not target: raise HTTPException(404, "User not found")
    if payload.role not in ("student", "teacher", "school_admin", "super_admin"):
        raise HTTPException(400, "Invalid role.")
    target.role = payload.role
    if payload.school_id is not None:
        target.school_id = payload.school_id
    db.commit()
    log_activity(db, "USER_ROLE_UPDATED", f"Admin set {target.display_name}'s role to {payload.role}")
    return {"ok": True, "role": target.role, "school_id": target.school_id}


# FIX: Changed to POST so the browser doesn't block the request body
@app.post("/admin/students/bulk-delete")
def admin_bulk_delete_students(payload: dict, db: Session = Depends(get_db), admin: models.User = Depends(require_super_admin)):
    user_ids = payload.get("user_ids", [])
    if not user_ids: return {"ok": True, "deleted": 0}
    
    for uid in user_ids:
        # 1. Delete Test Answers
        test_ids_subq = db.query(models.Test.id).filter(models.Test.user_id == uid).subquery()
        db.query(models.TestAnswer).filter(models.TestAnswer.test_id.in_(test_ids_subq)).delete(synchronize_session=False)
        
        # 2. Delete Tests
        db.query(models.Test).filter(models.Test.user_id == uid).delete(synchronize_session=False)
        
        # 3. Delete Bookmarks
        db.query(models.QuestionBookmark).filter(models.QuestionBookmark.user_id == uid).delete(synchronize_session=False)
        
        # 4. Delete Emails
        db.query(models.EmailLog).filter(models.EmailLog.user_id == uid).delete(synchronize_session=False)
        
        # 5. Delete Daily Challenge Logs (This caused the crash!)
        db.query(models.DailyChallengeLog).filter(models.DailyChallengeLog.user_id == uid).delete(synchronize_session=False)
        
        # 6. Delete Reported Questions made by the user
        db.query(models.ReportedQuestion).filter(models.ReportedQuestion.user_id == uid).delete(synchronize_session=False)
        
        # 7. Delete Inventory (Vault items)
        db.query(models.UserInventory).filter(models.UserInventory.user_id == uid).delete(synchronize_session=False)
        
        # 8. Delete Study Sessions
        db.query(models.StudySession).filter(models.StudySession.user_id == uid).delete(synchronize_session=False)
        
        # 9. Delete Squad Membership
        db.query(models.SquadMember).filter(models.SquadMember.user_id == uid).delete(synchronize_session=False)
        
        # 10. Finally, Delete the User
        db.query(models.User).filter(models.User.id == uid).delete(synchronize_session=False)
        
    db.commit()
    log_activity(db, "STUDENTS_BULK_DELETED", f"Admin deleted {len(user_ids)} students")
    return {"ok": True, "deleted": len(user_ids)}


@app.get("/admin/students/{user_id}")
def admin_student_profile(user_id: int, db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    s = _scoped_user_or_403(db, admin, user_id)
    tests = db.query(models.Test).filter(models.Test.user_id == s.id).order_by(models.Test.started_at.desc()).all()
    history = [{
        "id": t.id, "subject": t.subject, "mode": t.mode,
        "question_count": t.question_count, "score": t.score, "time_taken_seconds": t.time_taken_seconds,
        "started_at": t.started_at.isoformat() if t.started_at else None,
        "submitted_at": t.submitted_at.isoformat() if t.submitted_at else None,
    } for t in tests]
    return {
        "user": {
            "id": s.id, "username": s.username, "display_name": s.display_name,
            "school_id": s.school_id, "school_name": s.school_name,
            "education_level": s.education_level, "exam_target": s.exam_target, "subjects": s.subjects,
            "avatar_id": s.avatar_id, "is_active": s.is_active, "email": s.email,
            "role": s.role, "xp": s.xp, "coins": s.coins, "is_premium": s.is_premium,
            "created_at": s.created_at.isoformat() if s.created_at else None,
        },
        "tests": history,
    }


@app.post("/admin/reset-password")
def admin_reset_password(payload: AdminResetPasswordRequest, db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    target = _scoped_user_or_403(db, admin, payload.user_id)
    target.password_hash = hash_password(payload.new_password)
    db.commit()
    return {"ok": True, "user_id": target.id}


@app.post("/admin/students/{user_id}/toggle-active")
def admin_toggle_active(user_id: int, db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    target = _scoped_user_or_403(db, admin, user_id)
    target.is_active = not target.is_active
    db.commit()
    return {"ok": True, "is_active": target.is_active}


@app.get("/admin/emails")
def admin_get_emails(db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    q = db.query(models.EmailLog)
    if is_school_scoped(admin):
        user_ids_subq = db.query(models.User.id).filter(models.User.school_id == admin.school_id).subquery()
        q = q.filter(models.EmailLog.user_id.in_(user_ids_subq))
    logs = q.order_by(models.EmailLog.sent_at.desc()).limit(50).all()
    return [{"id": l.id, "recipient": l.recipient, "subject": l.subject, "body": l.body, "reply": l.reply, "is_read": l.is_read, "sent_at": l.sent_at.isoformat()} for l in logs]


@app.get("/admin/emails/unread")
def admin_unread_count(db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    q = db.query(models.EmailLog).filter(models.EmailLog.is_read == False)
    if is_school_scoped(admin):
        user_ids_subq = db.query(models.User.id).filter(models.User.school_id == admin.school_id).subquery()
        q = q.filter(models.EmailLog.user_id.in_(user_ids_subq))
    return {"unread": q.count()}


@app.put("/admin/emails/{email_id}/read")
def admin_mark_read(email_id: int, db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    log = db.query(models.EmailLog).filter(models.EmailLog.id == email_id).first()
    if not log: raise HTTPException(404, "Email not found")
    if is_school_scoped(admin):
        if not log.user_id: raise HTTPException(403, "This email is not linked to a user in your school.")
        owner = db.query(models.User).filter(models.User.id == log.user_id).first()
        if not owner or owner.school_id != admin.school_id:
            raise HTTPException(403, "This email belongs to another school.")
    log.is_read = True; db.commit()
    return {"ok": True}


@app.post("/admin/emails/{email_id}/reply")
async def admin_reply_email(email_id: int, payload: dict, background_tasks: BackgroundTasks, db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    log = db.query(models.EmailLog).filter(models.EmailLog.id == email_id).first()
    if not log: raise HTTPException(404, "Email not found")
    if is_school_scoped(admin):
        if not log.user_id: raise HTTPException(403, "This email is not linked to a user in your school.")
        owner = db.query(models.User).filter(models.User.id == log.user_id).first()
        if not owner or owner.school_id != admin.school_id:
            raise HTTPException(403, "This email belongs to another school.")
    reply_text = payload.get("message", "")
    if not reply_text.strip(): raise HTTPException(400, "Reply cannot be empty")
    log.reply = reply_text
    log.replied_at = datetime.now(timezone.utc)
    db.commit()
    background_tasks.add_task(email_service.send_support_email, "Admin Reply", log.recipient, reply_text)
    return {"ok": True, "message": "Reply sent"}


@app.get("/zoom-sessions")
def public_zoom(db: Session = Depends(get_db)):
    now = datetime.now(timezone.utc)
    rows = db.query(models.ZoomSession).filter(models.ZoomSession.scheduled_at >= now).order_by(models.ZoomSession.scheduled_at.asc()).all()
    return [{"id": r.id, "title": r.title, "subject": r.subject, "scheduled_at": r.scheduled_at.isoformat(), "zoom_url": r.zoom_url, "description": r.description} for r in rows]


@app.get("/admin/zoom-sessions")
def admin_list_zoom(db: Session = Depends(get_db), admin: models.User = Depends(require_admin)):
    rows = db.query(models.ZoomSession).order_by(models.ZoomSession.scheduled_at.desc()).all()
    return [{"id": r.id, "title": r.title, "subject": r.subject, "scheduled_at": r.scheduled_at.isoformat(), "zoom_url": r.zoom_url, "description": r.description} for r in rows]


@app.post("/admin/zoom-sessions", response_model=ZoomSessionOut)
def admin_create_zoom(payload: ZoomSessionUpsert, db: Session = Depends(get_db), admin: models.User = Depends(require_super_admin)):
    try: when = datetime.fromisoformat(payload.scheduled_at.replace("Z", "+00:00"))
    except Exception: raise HTTPException(400, "scheduled_at must be ISO format")
    z = models.ZoomSession(title=payload.title, subject=payload.subject, scheduled_at=when, zoom_url=payload.zoom_url, description=payload.description)
    db.add(z); db.commit(); db.refresh(z)
    return ZoomSessionOut(id=z.id, title=z.title, subject=z.subject, scheduled_at=z.scheduled_at.isoformat(), zoom_url=z.zoom_url, description=z.description)


@app.delete("/admin/zoom-sessions/{zid}")
def admin_delete_zoom(zid: int, db: Session = Depends(get_db), admin: models.User = Depends(require_super_admin)):
    z = db.query(models.ZoomSession).filter(models.ZoomSession.id == zid).first()
    if not z: raise HTTPException(404, "zoom session not found")
    db.delete(z); db.commit()
    return {"ok": True}


@app.post("/ai/explain")
def explain_question(payload: dict, db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    check_trial_access(current_user)
    if not check_rate_limit(current_user.id):
        raise HTTPException(status_code=429, detail="You have reached your daily limit of 5 AI explanations. Please try again tomorrow.")
    question_text = payload.get("question_text")
    correct_answer = payload.get("correct_answer")
    if not question_text or not correct_answer:
        raise HTTPException(status_code=400, detail="Missing question text or correct answer.")
    explanation = get_ai_explanation(question_text, correct_answer)
    return {"explanation": explanation}


@app.get("/subjects/mine")
def my_subjects(user: models.User = Depends(get_current_user)):
    return {"education_level": user.education_level, "exam_target": user.exam_target, "subjects": user.subjects}


# ---------------- DAILY CHALLENGE ----------------

@app.get("/daily-challenge")
def get_daily_challenge(db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    check_trial_access(current_user)
    today = datetime.now().strftime("%Y-%m-%d")
    log = db.query(DailyChallengeLog).filter(
        DailyChallengeLog.user_id == current_user.id,
        DailyChallengeLog.challenge_date == today
    ).first()
    if log:
        return {"status": "completed", "is_correct": log.is_correct}
    questions = db.query(Question).filter(Question.is_active == True).order_by(func.random()).limit(3).all()
    if len(questions) < 3:
        return {"status": "error", "message": "Not enough questions in the bank."}
    return {
        "status": "active",
        "questions": [{
            "id": q.id, "subject": q.subject, "question_text": q.question_text, "image_url": q.image_url,
            "options": [{"letter": "A", "text": q.option_a}, {"letter": "B", "text": q.option_b}, {"letter": "C", "text": q.option_c}, {"letter": "D", "text": q.option_d}]
        } for q in questions]
    }


@app.post("/daily-challenge/submit")
def submit_daily_challenge(payload: dict, db: Session = Depends(get_db), current_user: models.User = Depends(get_current_user)):
    check_trial_access(current_user)
    today = datetime.now().strftime("%Y-%m-%d")
    log = db.query(DailyChallengeLog).filter(
        DailyChallengeLog.user_id == current_user.id,
        DailyChallengeLog.challenge_date == today
    ).first()
    if log:
        raise HTTPException(status_code=400, detail="You already completed today's challenge.")
    answers = payload.get("answers", [])
    correct_count = 0
    results = []
    for ans in answers:
        q_id = ans.get("question_id")
        selected = ans.get("selected_option")
        question = db.query(Question).filter(Question.id == q_id).first()
        if question:
            is_correct = (question.correct_option == selected)
            if is_correct: correct_count += 1
            results.append({"question_id": q_id, "is_correct": is_correct, "correct_option": question.correct_option, "explanation": question.explanation})
    new_log = DailyChallengeLog(user_id=current_user.id, challenge_date=today, is_correct=(correct_count >= 2))
    db.add(new_log)
    xp_earned = 50 if correct_count >= 2 else 0
    coins_earned = 20 if correct_count >= 2 else 0
    if xp_earned > 0:
        current_user.xp += xp_earned
        current_user.coins += coins_earned
    db.commit()
    return {"total_correct": correct_count, "total_questions": len(answers), "results": results, "xp_earned": xp_earned, "coins_earned": coins_earned}


# ---------------- VAULT (SHOP) ----------------

VAULT_ITEMS = [
    {"id": "avatar_hacker", "name": "Neon Hacker", "type": "avatar", "price": 100, "icon": "👨‍💻", "desc": "Look like a cyber security genius."},
    {"id": "avatar_robot", "name": "Quantum Robot", "type": "avatar", "price": 200, "icon": "🤖", "desc": "A highly advanced AI robot."},
    {"id": "avatar_alien", "name": "Cosmic Alien", "type": "avatar", "price": 300, "icon": "👽", "desc": "A visitor from another galaxy."},
    {"id": "avatar_ninja", "name": "Cyber Ninja", "type": "avatar", "price": 400, "icon": "🥷", "desc": "A stealthy warrior from the future."},
    {"id": "avatar_samurai", "name": "Star Samurai", "type": "avatar", "price": 400, "icon": "⚔️", "desc": "A master of the plasma blade."},
    {"id": "avatar_pilot", "name": "Starfighter Pilot", "type": "avatar", "price": 250, "icon": "🚀", "desc": "Top gun of the galaxy."},
    {"id": "avatar_astronaut", "name": "Deep Space Explorer", "type": "avatar", "price": 300, "icon": "👨‍🚀", "desc": "Explorer of the unknown."},
    {"id": "avatar_doctor", "name": "Quantum Doctor", "type": "avatar", "price": 350, "icon": "👨‍⚕️", "desc": "Healer of the future."},
    {"id": "avatar_scientist", "name": "Mad Scientist", "type": "avatar", "price": 350, "icon": "👨‍🔬", "desc": "Genius is often misunderstood."},
    {"id": "avatar_wizard", "name": "Tech Wizard", "type": "avatar", "price": 500, "icon": "🧙‍♂️", "desc": "Coding spells and casting algorithms."},
    {"id": "avatar_vampire", "name": "Neon Vampire", "type": "avatar", "price": 450, "icon": "🧛", "desc": "Child of the night."},
    {"id": "avatar_werewolf", "name": "Cyber Werewolf", "type": "avatar", "price": 450, "icon": "🐺", "desc": "Half beast, half machine."},
    {"id": "avatar_gorilla", "name": "Gorilla Warrior", "type": "avatar", "price": 500, "icon": "🦍", "desc": "Strength and honor."},
    {"id": "avatar_dragon", "name": "Space Dragon", "type": "avatar", "price": 800, "icon": "🐉", "desc": "Fire breather of the cosmos."},
    {"id": "avatar_tiger", "name": "Saber Tiger", "type": "avatar", "price": 600, "icon": "🐯", "desc": "Ferocious and fast."},
    {"id": "avatar_lion", "name": "Alpha Lion", "type": "avatar", "price": 600, "icon": "🦁", "desc": "King of the concrete jungle."},
    {"id": "avatar_eagle", "name": "Sky Eagle", "type": "avatar", "price": 500, "icon": "🦅", "desc": "Eyes of a hunter."},
    {"id": "avatar_phoenix", "name": "Phoenix Rising", "type": "avatar", "price": 1000, "icon": "🔥", "desc": "Reborn from the ashes."},
    {"id": "avatar_ghost", "name": "Ghost Protocol", "type": "avatar", "price": 500, "icon": "👻", "desc": "You can't hit what you can't see."},
    {"id": "avatar_king", "name": "The Sovereign", "type": "avatar", "price": 1500, "icon": "👑", "desc": "Bow before the throne."},
    {"id": "title_scavenger", "name": "The Scavenger", "type": "title", "price": 50, "icon": "🥾", "desc": "Everyone starts somewhere."},
    {"id": "title_survivor", "name": "The Survivor", "type": "title", "price": 150, "icon": "🩹", "desc": "You've been through the wars."},
    {"id": "title_student", "name": "Eternal Student", "type": "title", "price": 200, "icon": "📚", "desc": "Always learning, always growing."},
    {"id": "title_gunslinger", "name": "The Gunslinger", "type": "title", "price": 300, "icon": "🔫", "desc": "Quick on the draw."},
    {"id": "title_scientist", "name": "The Scientist", "type": "title", "price": 400, "icon": "🔬", "desc": "Knowledge is power."},
    {"id": "title_thinker", "name": "The Thinker", "type": "title", "price": 450, "icon": "🧠", "desc": "Outsmarting the opposition."},
    {"id": "title_analyst", "name": "The Analyst", "type": "title", "price": 450, "icon": "📊", "desc": "Data is your weapon."},
    {"id": "title_strategist", "name": "The Strategist", "type": "title", "price": 500, "icon": "♟️", "desc": "Always ten steps ahead."},
    {"id": "title_destroyer", "name": "JAMB Destroyer", "type": "title", "price": 600, "icon": "⚔️", "desc": "Show everyone you crush exams."},
    {"id": "title_genius", "name": "Certified Genius", "type": "title", "price": 750, "icon": "🧠", "desc": "For the absolute top scorers."},
    {"id": "title_overlord", "name": "The Overlord", "type": "title", "price": 800, "icon": "😈", "desc": "Ruler of the leaderboard."},
    {"id": "title_sage", "name": "The Sage", "type": "title", "price": 850, "icon": "🧘", "desc": "Wisdom beyond your years."},
    {"id": "title_ghost", "name": "The Ghost", "type": "title", "price": 900, "icon": "👻", "desc": "No one sees you coming."},
    {"id": "title_lord", "name": "Lord Makaverse", "type": "title", "price": 1000, "icon": "👑", "desc": "The ultimate flex."},
    {"id": "title_legend", "name": "Living Legend", "type": "title", "price": 2000, "icon": "🌟", "desc": "Your name will be remembered."},
    {"id": "powerup_ai_pass", "name": "AI Unlimited Pass", "type": "powerup", "price": 300, "icon": "🤖", "desc": "Unlimited AI explanations for 24 hours."},
    {"id": "powerup_streak_shield", "name": "Streak Shield", "type": "powerup", "price": 400, "icon": "🛡️", "desc": "Protects your streak for one missed day."},
    {"id": "powerup_retake_token", "name": "Mock Retake Token", "type": "powerup", "price": 600, "icon": "🎟️", "desc": "Allows one retake of a failed Mock Exam."},
    {"id": "powerup_xp_boost", "name": "XP Boost (2x)", "type": "powerup", "price": 500, "icon": "⚡", "desc": "Double XP for 1 hour."},
    {"id": "powerup_coin_boost", "name": "Coin Boost (2x)", "type": "powerup", "price": 500, "icon": "🪙", "desc": "Double Coins for 1 hour."},
    {"id": "powerup_hint", "name": "Question Hint", "type": "powerup", "price": 150, "icon": "💡", "desc": "Get a hint for one question."},
    {"id": "powerup_skip", "name": "Skip Question", "type": "powerup", "price": 200, "icon": "⏭️", "desc": "Skip a question without penalty."},
    {"id": "powerup_mystery_box", "name": "Mystery Box", "type": "powerup", "price": 250, "icon": "🎁", "desc": "Contains a random prize (XP, Coins, or a Power-Up)."},
    {"id": "powerup_title_color", "name": "Custom Title Color", "type": "powerup", "price": 350, "icon": "🎨", "desc": "Make your title stand out on the leaderboard."},
    {"id": "powerup_avatar_frame", "name": "Golden Avatar Frame", "type": "powerup", "price": 450, "icon": "🖼️", "desc": "A golden frame for your profile picture."},
    {"id": "powerup_extra_life", "name": "Extra Life", "type": "powerup", "price": 700, "icon": "❤️", "desc": "Prevents one auto-submit from tab switching."},
    {"id": "powerup_double_coins", "name": "Double Coins (Daily)", "type": "powerup", "price": 800, "icon": "💰", "desc": "Doubles the coins you earn from tests for a day."},
    {"id": "powerup_double_xp", "name": "Double XP (Daily)", "type": "powerup", "price": 800, "icon": "🌟", "desc": "Doubles the XP you earn from tests for a day."},
    {"id": "powerup_library_access", "name": "Library Access (7 Days)", "type": "powerup", "price": 500, "icon": "📚", "desc": "Temporary access to the Study Library."},
    {"id": "powerup_premium_trial", "name": "Premium Trial (3 Days)", "type": "powerup", "price": 1000, "icon": "✨", "desc": "Taste the Premium life for 3 days."},
]


@app.get("/vault/items")
def get_vault_items(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    check_trial_access(user)
    return VAULT_ITEMS


@app.get("/vault/inventory")
def get_my_inventory(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    items = db.query(UserInventory).filter(UserInventory.user_id == user.id).all()
    return [{"item_id": i.item_id, "item_type": i.item_type, "purchased_at": i.purchased_at.isoformat()} for i in items]


@app.post("/vault/buy")
def buy_vault_item(payload: dict, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    check_trial_access(user)
    item_id = payload.get("item_id")
    item = next((i for i in VAULT_ITEMS if i["id"] == item_id), None)
    if not item: raise HTTPException(400, "Item not found")
    if user.coins < item["price"]: raise HTTPException(400, "Not enough coins!")
    if item["type"] in ["avatar", "title"]:
        existing = db.query(UserInventory).filter(UserInventory.user_id == user.id, UserInventory.item_id == item_id).first()
        if existing: raise HTTPException(400, "You already own this item.")
    elif item["type"] == "powerup":
        last_purchase = db.query(UserInventory).filter(UserInventory.user_id == user.id, UserInventory.item_id == item_id).order_by(UserInventory.purchased_at.desc()).first()
        if last_purchase:
            days_since = (datetime.now(timezone.utc) - last_purchase.purchased_at).days
            if days_since < 7:
                raise HTTPException(400, f"You can buy this again in {7 - days_since} days.")
    user.coins -= item["price"]
    new_item = UserInventory(user_id=user.id, item_id=item["id"], item_type=item["type"])
    db.add(new_item)
    if item["type"] == "avatar": user.avatar_id = item["id"]
    elif item["type"] == "title": user.active_title = item["name"]
    db.commit()
    log_activity(db, "VAULT_PURCHASE", f"{user.display_name} bought {item['name']}")
    return {"ok": True, "message": "Purchase successful!", "coins": user.coins}


@app.post("/vault/equip")
def equip_vault_item(payload: dict, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    item_id = payload.get("item_id")
    item = next((i for i in VAULT_ITEMS if i["id"] == item_id), None)
    if not item: raise HTTPException(400, "Item not found")
    owns_item = db.query(UserInventory).filter(UserInventory.user_id == user.id, UserInventory.item_id == item_id).first()
    if not owns_item: raise HTTPException(400, "You don't own this item.")
    if item["type"] == "avatar": user.avatar_id = item["id"]
    elif item["type"] == "title": user.active_title = item["name"]
    db.commit()
    return {"ok": True, "message": "Equipped successfully!"}


# ---------------- POWER LADDER ----------------

RANKS = [
    {"level": 1, "name": "Cadet", "req": 0}, {"level": 2, "name": "Scout", "req": 500}, {"level": 3, "name": "Ranger", "req": 1200},
    {"level": 4, "name": "Pilot", "req": 2100}, {"level": 5, "name": "Ensign", "req": 3200}, {"level": 6, "name": "Lieutenant", "req": 4500},
    {"level": 7, "name": "Captain", "req": 6000}, {"level": 8, "name": "Commander", "req": 7700}, {"level": 9, "name": "Major", "req": 9600},
    {"level": 10, "name": "Colonel", "req": 11700}, {"level": 11, "name": "Vanguard", "req": 14000}, {"level": 12, "name": "Centurion", "req": 16500},
    {"level": 13, "name": "Starfleet", "req": 19200}, {"level": 14, "name": "Sentinel", "req": 22100}, {"level": 15, "name": "Overlord", "req": 25200},
    {"level": 16, "name": "Archon", "req": 28500}, {"level": 17, "name": "Sovereign", "req": 32000}, {"level": 18, "name": "Ascendant", "req": 35700},
    {"level": 19, "name": "Transcendent", "req": 39600}, {"level": 20, "name": "Lord Makaverse", "req": 43700},
]


@app.get("/ranks")
def get_ranks(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    return {"user_xp": user.xp, "ranks": RANKS}


# ---------------- STUDY TIME ----------------

@app.post("/study-time")
def save_study_time(payload: dict, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    seconds = payload.get("seconds", 0)
    if seconds > 0:
        user.total_study_time += seconds
        db.commit()
    return {"ok": True, "total": user.total_study_time}


# ---------------- PAYSTACK ----------------

PAYSTACK_BASE = "https://api.paystack.co"


def _paystack_secret() -> str:
    key = os.getenv("PAYSTACK_SECRET_KEY", "").strip()
    if not key:
        raise HTTPException(503, "Payments are not configured. Please set PAYSTACK_SECRET_KEY.")
    return key


@app.post("/payments/initialize")
def init_payment(payload: InitPaymentRequest, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    secret = _paystack_secret()
    if payload.plan_type != "monthly": raise HTTPException(400, "Only the monthly plan is available.")
    amount_kobo = 500000
    if not user.email: raise HTTPException(400, "Please add an email to your profile before upgrading.")
    reference = f"MK_{user.id}_{secrets.token_hex(8)}"
    txn = Transaction(user_id=user.id, reference=reference, amount=amount_kobo / 100.0, plan_type="monthly", status="pending")
    db.add(txn); db.commit()
    payload_body = {
        "email": user.email, "amount": amount_kobo, "reference": reference,
        "callback_url": f"{os.getenv('FRONTEND_URL', 'http://localhost:5173')}/upgrade",
        "metadata": {"user_id": user.id, "plan_type": "monthly", "display_name": user.display_name},
    }
    try:
        with httpx.Client(timeout=20.0) as client:
            resp = client.post(f"{PAYSTACK_BASE}/transaction/initialize", json=payload_body,
                headers={"Authorization": f"Bearer {secret}", "Content-Type": "application/json"})
    except httpx.HTTPError as e:
        db.query(Transaction).filter(Transaction.reference == reference).delete(synchronize_session=False); db.commit()
        raise HTTPException(502, f"Could not reach Paystack: {str(e)}")
    data = resp.json() if resp.content else {}
    if resp.status_code >= 400 or not data.get("status"):
        db.query(Transaction).filter(Transaction.reference == reference).delete(synchronize_session=False); db.commit()
        raise HTTPException(502, data.get("message") or f"Paystack error {resp.status_code}")
    auth_data = data.get("data") or {}
    return {"ok": True, "authorization_url": auth_data.get("authorization_url"), "access_code": auth_data.get("access_code"), "reference": reference}


@app.post("/payments/webhook")
async def paystack_webhook(request: Request, db: Session = Depends(get_db)):
    secret = os.getenv("PAYSTACK_SECRET_KEY", "").strip()
    if not secret: raise HTTPException(503, "Webhook not configured")
    raw = await request.body()
    signature = request.headers.get("x-paystack-signature", "")
    expected = hmac.new(secret.encode("utf-8"), raw, hashlib.sha512).hexdigest()
    if not hmac.compare_digest(expected, signature): raise HTTPException(401, "Invalid signature")
    try: event = __import__("json").loads(raw.decode("utf-8"))
    except Exception: raise HTTPException(400, "Invalid JSON")
    if event.get("event") != "charge.success": return {"ok": True, "ignored": True}
    data = event.get("data") or {}
    reference = data.get("reference")
    if not reference: return {"ok": True, "ignored": True}
    txn = db.query(Transaction).filter(Transaction.reference == reference).first()
    if not txn: return {"ok": True, "ignored": True}
    if txn.status == "success": return {"ok": True, "already_processed": True}
    if txn.plan_type != "monthly": return {"ok": True, "ignored": True}
    user = db.query(models.User).filter(models.User.id == txn.user_id).first()
    if not user: return {"ok": True, "ignored": True}
    now = datetime.now(timezone.utc); expires = now + timedelta(days=30)
    txn.status = "success"; txn.paid_at = now; txn.expires_at = expires
    user.is_premium = True; user.premium_expires_at = expires
    db.commit()
    log_activity(db, "PAYMENT_SUCCESS", f"{user.display_name} upgraded to monthly ({reference})")
    return {"ok": True}


@app.post("/payments/verify")
async def verify_payment(payload: dict, db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    reference = payload.get("reference")
    if not reference: raise HTTPException(400, "Missing payment reference")
    secret = _paystack_secret()
    try:
        with httpx.Client(timeout=20.0) as client:
            resp = client.get(f"{PAYSTACK_BASE}/transaction/verify/{reference}", headers={"Authorization": f"Bearer {secret}"})
    except httpx.HTTPError as e:
        raise HTTPException(502, f"Could not reach Paystack: {str(e)}")
    data = resp.json()
    if resp.status_code >= 400 or not data.get("status"): raise HTTPException(400, "Invalid reference")
    paystack_data = data.get("data", {})
    if paystack_data.get("status") != "success": raise HTTPException(400, "Payment was not successful")
    txn = db.query(Transaction).filter(Transaction.reference == reference).first()
    if not txn: raise HTTPException(404, "Transaction not found in our records")
    if txn.status == "success": return {"ok": True, "message": "Already verified"}
    now = datetime.now(timezone.utc); expires = now + timedelta(days=30)
    txn.status = "success"; txn.paid_at = now; txn.expires_at = expires
    user.is_premium = True; user.premium_expires_at = expires
    db.commit()
    log_activity(db, "PAYMENT_VERIFIED", f"{user.display_name} verified payment {reference} manually")
    return {"ok": True, "message": "Payment verified and account upgraded"}


@app.get("/payments/history")
def get_payment_history(db: Session = Depends(get_db), user: models.User = Depends(get_current_user)):
    txns = db.query(Transaction).filter(Transaction.user_id == user.id).order_by(Transaction.created_at.desc()).limit(3).all()
    return [{"id": t.id, "reference": t.reference, "amount": t.amount, "status": t.status,
             "paid_at": t.paid_at.isoformat() if t.paid_at else None,
             "expires_at": t.expires_at.isoformat() if t.expires_at else None} for t in txns]


# ---------------- ADMIN: VAULT & PREMIUM (super_admin only) ----------------

@app.get("/admin/vault/purchases")
def admin_get_vault_purchases(db: Session = Depends(get_db), admin: models.User = Depends(require_super_admin)):
    purchases = db.query(UserInventory).order_by(UserInventory.purchased_at.desc()).limit(100).all()
    users = {u.id: u for u in db.query(models.User).filter(models.User.id.in_([p.user_id for p in purchases])).all()}
    return [{"id": p.id, "user_id": p.user_id,
             "display_name": users[p.user_id].display_name if p.user_id in users else "Unknown",
             "item_id": p.item_id, "item_type": p.item_type, "purchased_at": p.purchased_at.isoformat()} for p in purchases]


@app.post("/admin/vault/grant")
def admin_grant_vault_item(payload: dict, db: Session = Depends(get_db), admin: models.User = Depends(require_super_admin)):
    user_id = payload.get("user_id"); item_id = payload.get("item_id")
    item = next((i for i in VAULT_ITEMS if i["id"] == item_id), None)
    if not item: raise HTTPException(400, "Item not found")
    target = db.query(models.User).filter(models.User.id == user_id).first()
    if not target: raise HTTPException(404, "User not found")
    new_item = UserInventory(user_id=user_id, item_id=item_id, item_type=item["type"])
    db.add(new_item); db.commit()
    log_activity(db, "VAULT_GRANTED", f"Admin granted {item['name']} to {target.display_name}")
    return {"ok": True}


@app.get("/admin/payments")
def admin_get_payments(db: Session = Depends(get_db), admin: models.User = Depends(require_super_admin)):
    txns = db.query(Transaction).order_by(Transaction.created_at.desc()).limit(100).all()
    users = {u.id: u for u in db.query(models.User).filter(models.User.id.in_([t.user_id for t in txns])).all()}
    return [{"id": t.id, "user_id": t.user_id,
             "display_name": users[t.user_id].display_name if t.user_id in users else "Unknown",
             "reference": t.reference, "amount": t.amount, "plan_type": t.plan_type, "status": t.status,
             "retry_count": t.retry_count,
             "paid_at": t.paid_at.isoformat() if t.paid_at else None,
             "expires_at": t.expires_at.isoformat() if t.expires_at else None} for t in txns]


@app.post("/admin/payments/{ref}/resolve")
def admin_resolve_payment(ref: str, payload: ResolvePaymentRequest, db: Session = Depends(get_db), admin: models.User = Depends(require_super_admin)):
    txn = db.query(Transaction).filter(Transaction.reference == ref).first()
    if not txn: raise HTTPException(404, "Transaction not found")
    user = db.query(models.User).filter(models.User.id == txn.user_id).first()
    if not user: raise HTTPException(404, "User not found")
    if payload.action == "success" and txn.status != "success":
        now = datetime.now(timezone.utc); expires = now + timedelta(days=30)
        txn.status = "success"; txn.paid_at = now; txn.expires_at = expires
        user.is_premium = True; user.premium_expires_at = expires
    elif payload.action == "failed":
        txn.status = "failed"
        txn.retry_count = (txn.retry_count or 0) + 1
    db.commit()
    log_activity(db, "PAYMENT_MANUAL_RESOLVE", f"Admin {payload.action}ed payment {ref}")
    return {"ok": True}


@app.post("/contact/enquiry")
def public_enquiry(payload: dict, background_tasks: BackgroundTasks, db: Session = Depends(get_db)):
    name = payload.get("name", "").strip()
    email = payload.get("email", "").strip()
    message = payload.get("message", "").strip()
    if not name or not email or not message: raise HTTPException(400, "Missing required fields")
    support_email = "makaverseedtech@mail.com"
    body_text = f"Name: {name}\nEmail: {email}\n\nMessage:\n{message}"
    background_tasks.add_task(email_service.send_support_email, name, email, body_text)
    log = models.EmailLog(user_id=None, recipient=support_email, subject="New Enquiry from " + name, body=message)
    db.add(log); db.commit()
    return {"ok": True}


@app.get("/verify-certificate/{cert_id}")
def verify_certificate(cert_id: str, db: Session = Depends(get_db)):
    parts = cert_id.split("-")
    if len(parts) != 3 or parts[0] != "MK": raise HTTPException(400, "Invalid Certificate ID format.")
    cert_type = parts[1].upper()
    try: test_id = int(parts[2])
    except ValueError: raise HTTPException(400, "Invalid Certificate ID format.")
    test = db.query(models.Test).filter(models.Test.id == test_id).first()
    if not test: raise HTTPException(404, "Certificate not found.")
    if cert_type == "MOCK" and test.mode != "mock": raise HTTPException(404, "Certificate not found (mock mismatch).")
    if cert_type == "CHAMP" and test.mode != "championship": raise HTTPException(404, "Certificate not found (championship mismatch).")
    pct = (test.score / test.question_count) * 100 if test.question_count > 0 else 0
    if pct < 40: raise HTTPException(404, "Certificate is invalid (score below pass mark).")
    user = db.query(models.User).filter(models.User.id == test.user_id).first()
    if not user: raise HTTPException(404, "Student not found.")
    return {"valid": True, "student_name": user.display_name, "school_name": user.school_name,
            "subject": test.subject, "score": test.score, "total_questions": test.question_count,
            "percentage": round(pct, 1), "date": test.submitted_at.isoformat() if test.submitted_at else None,
            "mode": test.mode}